"""Integration interfaces per channel, loaded from Excel or managed by hand, checked against the systems catalogue."""

import io
import re
import uuid
from datetime import datetime, timezone

from app import db
from app.config import kb_countries, kb_interface_channels
from app.errors import AppError
from app.permissions import require
from app.services import gcs, systems

PROTOCOLS = ("API", "MQ", "Kafka", "gRPC", "FTP", "SMTP", "SMPP")
HEADERS = ("Interface name", "Endpoint", "Description", "Protocol", "Source", "Destination")
MAX_ROWS = 5000
MAX_BYTES = 20 * 1024 * 1024
MAX_NAME = 200
MAX_ENDPOINT = 500
MAX_DESCRIPTION = 2000
MAX_RULES = 20000
MAX_FIELDS = 200
MAX_PATH = 200
MAX_FIELD_LENGTH = 200
MAX_FIELD_VALIDATION = 1000
MAX_EXAMPLE = 500
EMBED_CHARS = 8000
FIELD_TYPES = ("string", "number", "integer", "boolean", "date", "datetime", "object", "array", "enum")
FIELD_DIRECTIONS = ("request", "response")
REQUIRED_VALUES = ("Y", "N", "C")
_REQUIRED_WORDS = {"Y": "required", "N": "optional", "C": "conditional"}


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _require_editor(user: dict) -> None:
    require(user, "knowledge", "write")


def _channel(value: str) -> str:
    channels = kb_interface_channels()
    channel = (value or "").strip().upper()
    if channel not in channels:
        raise AppError(404, "Unknown interface channel")
    return channel


def _canonical(value: str, allowed: tuple[str, ...]) -> str | None:
    lookup = {item.lower(): item for item in allowed}
    return lookup.get((value or "").strip().lower())


def country_label(value: str) -> str:
    """Empty means the interface applies in every catalog country."""
    codes = [part.strip() for part in (value or "").split(",") if part.strip()]
    return ", ".join(codes) if codes else "All countries"


def _country(value: str) -> str:
    """Blank is all countries. Several codes are stored in catalog order, comma-separated."""
    text = (value or "").strip()
    if not text or text.lower() in {"all", "all countries"}:
        return ""
    allowed = kb_countries()
    order = {code: index for index, code in enumerate(allowed)}
    chosen: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[\s,]+", text):
        if not part:
            continue
        canonical = _canonical(part, allowed)
        if canonical is None:
            raise AppError(400, f"Country must be one of: {', '.join(allowed)}.")
        if canonical not in seen:
            seen.add(canonical)
            chosen.append(canonical)
    chosen.sort(key=lambda code: order[code])
    return ",".join(chosen)


def _protocol(value: str) -> str:
    canonical = _canonical(value, PROTOCOLS)
    if canonical is None:
        raise AppError(400, f"Protocol must be one of: {', '.join(PROTOCOLS)}.")
    return canonical


def _system_codes(value: str, label: str) -> str:
    raw = (value or "").replace(";", ",")
    codes = systems.split_codes(raw)
    if not codes:
        raise AppError(400, f"{label} is required.")
    known = systems.existing_codes()
    unknown = [code for code in codes if code not in known]
    if unknown:
        raise AppError(400, f"{label} '{', '.join(unknown)}' is not a known system code.")
    return ", ".join(codes)


def _endpoint(value: str) -> str:
    text = (value or "").strip()
    if len(text) > MAX_ENDPOINT:
        raise AppError(400, f"Endpoint must be {MAX_ENDPOINT} characters or fewer.")
    return text


def _limit(value: str, limit: int, label: str) -> str:
    text = (value or "").strip()
    if len(text) > limit:
        raise AppError(400, f"{label} must be {limit} characters or fewer.")
    return text


def _rules(value: str, label: str) -> str:
    return _limit(value, MAX_RULES, label)


def _field_sentence(field: dict) -> str:
    """One field line for chat: request amount (number, required): validation. description."""
    required = _REQUIRED_WORDS.get(field.get("required") or "", "optional")
    line = f"{field.get('direction') or 'request'} {field.get('field_path') or ''} ({field.get('field_type') or 'string'}, {required})"
    validation = (field.get("validation") or "").strip()
    description = (field.get("description") or "").strip()
    if validation:
        line = f"{line}: {validation}"
    if description:
        line = f"{line}. {description}"
    return line


def text_for(row: dict) -> str:
    """The interface as one readable sentence: what is embedded, searched and shown to the chat model.

    Fields come before the long rule texts so a truncated embed still keeps the field list.
    """
    description = f" {row['description']}" if row.get("description") else ""
    endpoint = f" Endpoint {row['endpoint']}." if row.get("endpoint") else ""
    lines = [
        (
            f"Interface {row['interface_name']} ({row['channel']} channel, country {country_label(row.get('country') or '')}).{description}{endpoint} "
            f"Protocol {row['protocol']}. Source {row['source']} -> Destination {row['destination']}."
        )
    ]
    for field in row.get("fields") or []:
        lines.append(_field_sentence(field))
    processing = (row.get("processing_rules") or "").strip()
    validation = (row.get("validation_rules") or "").strip()
    if processing:
        lines.append(f"Processing rules: {processing}")
    if validation:
        lines.append(f"Validation rules: {validation}")
    return "\n".join(lines)


def fields_by_interface(interface_ids: list[str]) -> dict[str, list[dict]]:
    """Field rows grouped by interface, request first, then position."""
    if not interface_ids:
        return {}
    marks = ", ".join("?" for _ in interface_ids)
    rows = db.fetchall(
        f"""
        SELECT field_id, interface_id, direction, position, field_path, field_type, required, length, validation, description, example
        FROM interface_fields
        WHERE interface_id IN ({marks})
        ORDER BY direction, position, field_path
        """,
        list(interface_ids),
    )
    grouped: dict[str, list[dict]] = {}
    for row in rows:
        grouped.setdefault(row["interface_id"], []).append({key: row[key] for key in row if key != "interface_id"})
    return grouped


def attach_fields(rows: list[dict]) -> list[dict]:
    """Copy each row and add its fields. Used when embedding or searching, not on the list payload."""
    grouped = fields_by_interface([row["interface_id"] for row in rows])
    attached = []
    for row in rows:
        copy = dict(row)
        copy["fields"] = grouped.get(row["interface_id"], [])
        attached.append(copy)
    return attached


def _public_field(row: dict) -> dict:
    return {
        "field_id": row["field_id"],
        "direction": row["direction"],
        "position": row["position"],
        "field_path": row["field_path"],
        "field_type": row["field_type"],
        "required": row["required"],
        "length": row.get("length") or "",
        "validation": row.get("validation") or "",
        "description": row.get("description") or "",
        "example": row.get("example") or "",
    }


def _parse_fields(body: dict) -> list[dict] | None:
    """None means the caller did not send fields, so the stored rows stay."""
    if "fields" not in body:
        return None
    raw = body.get("fields")
    if not isinstance(raw, list):
        raise AppError(400, "Fields must be a list.")
    counts = {direction: 0 for direction in FIELD_DIRECTIONS}
    seen: set[tuple[str, str]] = set()
    parsed: list[dict] = []
    positions = {direction: 0 for direction in FIELD_DIRECTIONS}
    for item in raw:
        if not isinstance(item, dict):
            raise AppError(400, "Each field must be an object.")
        direction = (item.get("direction") or "").strip().lower()
        if direction not in FIELD_DIRECTIONS:
            raise AppError(400, "Field direction must be request or response.")
        path = _limit(str(item.get("field_path") or ""), MAX_PATH, "Field path")
        if not path:
            raise AppError(400, "Field path is required.")
        key = (direction, path)
        if key in seen:
            label = "Request" if direction == "request" else "Response"
            raise AppError(400, f"{label} field '{path}' is listed more than once.")
        seen.add(key)
        counts[direction] += 1
        if counts[direction] > MAX_FIELDS:
            label = "Request" if direction == "request" else "Response"
            raise AppError(400, f"{label} has more than {MAX_FIELDS} fields.")
        field_type = _canonical(str(item.get("field_type") or ""), FIELD_TYPES)
        if field_type is None:
            raise AppError(400, f"Field type must be one of: {', '.join(FIELD_TYPES)}.")
        required = (item.get("required") or "").strip().upper()
        if required not in REQUIRED_VALUES:
            raise AppError(400, "Required must be Y, N, or C.")
        parsed.append(
            {
                "direction": direction,
                "position": positions[direction],
                "field_path": path,
                "field_type": field_type,
                "required": required,
                "length": _limit(str(item.get("length") or ""), MAX_FIELD_LENGTH, "Length"),
                "validation": _limit(str(item.get("validation") or ""), MAX_FIELD_VALIDATION, "Validation"),
                "description": _limit(str(item.get("description") or ""), MAX_DESCRIPTION, "Field description"),
                "example": _limit(str(item.get("example") or ""), MAX_EXAMPLE, "Example"),
            }
        )
        positions[direction] += 1
    return parsed


def _replace_fields(interface_id: str, fields: list[dict]) -> None:
    db.exec_sql("DELETE FROM interface_fields WHERE interface_id = ?", [interface_id])
    for field in fields:
        db.exec_sql(
            """
            INSERT INTO interface_fields (
                field_id, interface_id, direction, position, field_path, field_type, required, length, validation, description, example
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                str(uuid.uuid4()),
                interface_id,
                field["direction"],
                field["position"],
                field["field_path"],
                field["field_type"],
                field["required"],
                field["length"],
                field["validation"],
                field["description"],
                field["example"],
            ],
        )


def _delete_interface_rows(interface_id: str) -> None:
    db.exec_sql("DELETE FROM interface_fields WHERE interface_id = ?", [interface_id])
    db.exec_sql("DELETE FROM kb_capability_interfaces WHERE interface_id = ?", [interface_id])
    db.exec_sql("DELETE FROM interfaces WHERE interface_id = ?", [interface_id])


def list_interfaces(channel: str) -> list[dict]:
    rows = db.fetchall(
        """
        SELECT interface_id, channel, country, interface_name, endpoint, description, protocol, source, destination
        FROM interfaces WHERE channel = ? ORDER BY created_at, interface_name
        """,
        [_channel(channel)],
    )
    return rows


def get_interface(channel: str, interface_id: str) -> dict:
    row = db.fetchone(
        """
        SELECT interface_id, channel, country, interface_name, endpoint, description, protocol, source, destination,
               processing_rules, validation_rules
        FROM interfaces WHERE channel = ? AND interface_id = ?
        """,
        [_channel(channel), interface_id],
    )
    if row is None:
        raise AppError(404, "Interface not found")
    detail = dict(row)
    detail["processing_rules"] = detail.get("processing_rules") or ""
    detail["validation_rules"] = detail.get("validation_rules") or ""
    detail["fields"] = [_public_field(field) for field in fields_by_interface([interface_id]).get(interface_id, [])]
    return detail


def create_interface(user: dict, channel: str, body: dict) -> dict:
    _require_editor(user)
    channel = _channel(channel)
    country = _country(str(body.get("country") or ""))
    name = (body.get("interface_name") or "").strip()
    if not name:
        raise AppError(400, "Interface name is required.")
    if len(name) > MAX_NAME:
        raise AppError(400, f"Interface name must be {MAX_NAME} characters or fewer.")
    description = (body.get("description") or "").strip()
    if len(description) > MAX_DESCRIPTION:
        raise AppError(400, f"Description must be {MAX_DESCRIPTION} characters or fewer.")
    endpoint = _endpoint(str(body.get("endpoint") or ""))
    protocol = _protocol(str(body.get("protocol") or ""))
    source = _system_codes(str(body.get("source") or ""), "Source")
    destination = _system_codes(str(body.get("destination") or ""), "Destination")
    processing_rules = _rules(str(body.get("processing_rules") or ""), "Processing rules")
    validation_rules = _rules(str(body.get("validation_rules") or ""), "Validation rules")
    fields = _parse_fields(body) or []
    interface_id = str(uuid.uuid4())
    db.exec_sql(
        """
        INSERT INTO interfaces (
            interface_id, channel, country, interface_name, endpoint, description, protocol, source, destination,
            processing_rules, validation_rules, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [interface_id, channel, country, name, endpoint, description, protocol, source, destination, processing_rules, validation_rules, _now()],
    )
    if fields:
        _replace_fields(interface_id, fields)
    from app.services import reindex

    reindex.embed_missing_in_background()
    gcs.upload_db()
    return get_interface(channel, interface_id)


def update_interface(user: dict, channel: str, interface_id: str, body: dict) -> dict:
    _require_editor(user)
    current = get_interface(channel, interface_id)
    country = _country(str(body.get("country", current["country"]) or ""))
    name = (body.get("interface_name", current["interface_name"]) or "").strip()
    if not name:
        raise AppError(400, "Interface name is required.")
    if len(name) > MAX_NAME:
        raise AppError(400, f"Interface name must be {MAX_NAME} characters or fewer.")
    description = (body.get("description", current["description"]) or "").strip()
    if len(description) > MAX_DESCRIPTION:
        raise AppError(400, f"Description must be {MAX_DESCRIPTION} characters or fewer.")
    endpoint = _endpoint(str(body.get("endpoint", current.get("endpoint") or "") or ""))
    protocol = _protocol(str(body.get("protocol", current["protocol"]) or ""))
    source = _system_codes(str(body.get("source", current["source"]) or ""), "Source")
    destination = _system_codes(str(body.get("destination", current["destination"]) or ""), "Destination")
    processing_rules = _rules(str(body.get("processing_rules", current.get("processing_rules") or "") or ""), "Processing rules")
    validation_rules = _rules(str(body.get("validation_rules", current.get("validation_rules") or "") or ""), "Validation rules")
    fields = _parse_fields(body)
    db.exec_sql(
        """
        UPDATE interfaces
        SET country = ?, interface_name = ?, endpoint = ?, description = ?, protocol = ?, source = ?, destination = ?,
            processing_rules = ?, validation_rules = ?, embedding = NULL
        WHERE interface_id = ?
        """,
        [country, name, endpoint, description, protocol, source, destination, processing_rules, validation_rules, interface_id],
    )
    if fields is not None:
        _replace_fields(interface_id, fields)
    from app.services import reindex

    reindex.embed_missing_in_background()
    gcs.upload_db()
    return get_interface(channel, interface_id)


def delete_interface(user: dict, channel: str, interface_id: str) -> None:
    _require_editor(user)
    get_interface(channel, interface_id)
    _delete_interface_rows(interface_id)
    gcs.upload_db()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppError(400, "Could not read the Excel file. Upload an .xlsx file.") from exc


def parse_and_validate(data: bytes) -> list[dict]:
    rows = _read_sheet(data)
    if not rows:
        raise AppError(400, "The Excel file is empty.")
    header = [value.strip().lower() for value in rows[0]]
    missing = [name for name in HEADERS if name.lower() not in header]
    if missing:
        raise AppError(400, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(HEADERS)}.")
    index = {name: header.index(name.lower()) for name in HEADERS}
    known = systems.existing_codes()
    records: list[dict] = []
    errors: list[dict] = []

    def error(row_number: int, field: str, message: str) -> None:
        errors.append({"row": row_number, "field": field, "message": message})

    for offset, row in enumerate(rows[1:], start=2):
        values = {name: row[index[name]] if index[name] < len(row) else "" for name in HEADERS}
        if not any(values.values()):
            continue
        record = {"description": values["Description"], "endpoint": values["Endpoint"].strip()}
        if len(record["endpoint"]) > MAX_ENDPOINT:
            error(offset, "Endpoint", f"Endpoint must be {MAX_ENDPOINT} characters or fewer.")
        if not values["Interface name"]:
            error(offset, "Interface name", "Interface name is required.")
        record["interface_name"] = values["Interface name"]
        protocol = _canonical(values["Protocol"], PROTOCOLS)
        if protocol is None:
            error(offset, "Protocol", f"'{values['Protocol']}' is not allowed. Use one of: {', '.join(PROTOCOLS)}.")
        record["protocol"] = protocol or ""
        for field, key in (("Source", "source"), ("Destination", "destination")):
            codes = systems.split_codes(values[field].replace(";", ","))
            unknown = [code for code in codes if code not in known]
            if not codes:
                error(offset, field, f"{field} is required.")
            elif unknown:
                error(offset, field, f"'{', '.join(unknown)}' is not a known system code.")
            record[key] = ", ".join(codes)
        records.append(record)
    if len(records) > MAX_ROWS:
        raise AppError(400, f"At most {MAX_ROWS} rows can be uploaded at once.")
    if errors:
        raise AppError(400, f"{len(errors)} problem(s) found. Nothing was imported.", errors)
    if not records:
        raise AppError(400, "The Excel file has no data rows.")
    return records


def replace_from_excel(user: dict, channel: str, country_raw: str, filename: str, data: bytes) -> list[dict]:
    _require_editor(user)
    channel = _channel(channel)
    country = _country(country_raw)
    if not (filename or "").lower().endswith(".xlsx"):
        raise AppError(400, "Upload an Excel file (.xlsx).")
    if len(data) > MAX_BYTES:
        raise AppError(400, "The file is larger than 20 MB.")
    records = parse_and_validate(data)
    now = _now()
    db.exec_sql("DELETE FROM interfaces WHERE channel = ? AND country = ?", [channel, country])
    for record in records:
        db.exec_sql(
            """
            INSERT INTO interfaces (
                interface_id, channel, country, interface_name, endpoint, description, protocol, source, destination, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                str(uuid.uuid4()),
                channel,
                country,
                record["interface_name"],
                record["endpoint"],
                record["description"],
                record["protocol"],
                record["source"],
                record["destination"],
                now,
            ],
        )
    from app.services import reindex

    reindex.embed_missing_in_background()
    gcs.upload_db()
    return list_interfaces(channel)
