"""The systems catalogue: a code, a name and a description."""

import io
from datetime import datetime, timezone

from app import db
from app.errors import AppError
from app.permissions import require
from app.services import gcs, llm

MAX_CODE = 50
MAX_NAME = 200
MAX_DESCRIPTION = 2000
MAX_UPLOAD_ROWS = 500
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
UPLOAD_HEADERS = ("SYSTEM_CODE", "SYSTEM_NAME", "DESCRIPTION")


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _require_editor(user: dict) -> None:
    require(user, "knowledge", "write")


def split_codes(value: str | None) -> list[str]:
    """A comma-separated list of system codes: trimmed, upper case, no blanks, no duplicates, order kept."""
    seen: dict[str, None] = {}
    for part in (value or "").split(","):
        code = part.strip().upper()
        if code:
            seen.setdefault(code, None)
    return list(seen)


def existing_codes() -> set[str]:
    return {row["system_code"] for row in db.fetchall("SELECT system_code FROM systems")}


def codes_with_batch() -> set[str]:
    return {row["system_code"] for row in db.fetchall("SELECT system_code FROM systems WHERE has_batch")}


def codes_with_report() -> set[str]:
    return {row["system_code"] for row in db.fetchall("SELECT system_code FROM systems WHERE has_report")}


def _flag(value, default: bool = False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in ("1", "true", "yes", "y", "on"):
        return True
    if text in ("0", "false", "no", "n", "off", ""):
        return False
    return default


def text_for(row: dict) -> str:
    """The system as one readable sentence: what is embedded, searched and shown to the chat model."""
    parts = [f"System {row['system_code']} - {row['system_name']}."]
    if row.get("description"):
        parts.append(f"{row['description']}.".replace("..", "."))
    return " ".join(parts)


def _embed(row: dict) -> None:
    vector = llm.embed(text_for(row)) if llm.embedding_configured() else None
    db.exec_sql("UPDATE systems SET embedding = ? WHERE system_code = ?", [vector, row["system_code"]])


def _payload(row: dict) -> dict:
    return {
        "system_code": row["system_code"],
        "system_name": row["system_name"],
        "description": row["description"],
        "upstreams": row["upstreams"],
        "downstreams": row["downstreams"],
        "has_batch": bool(row.get("has_batch")),
        "has_report": bool(row.get("has_report")),
        "has_capability": bool(row.get("has_capability")),
        "created_at": row["created_at"],
        "last_modified": row["last_modified"],
    }


def list_systems() -> list[dict]:
    return [_payload(row) for row in db.fetchall("SELECT * FROM systems ORDER BY system_code")]


def _row(code: str) -> dict:
    row = db.fetchone("SELECT * FROM systems WHERE system_code = ?", [code.strip().upper()])
    if row is None:
        raise AppError(404, "System not found")
    return row


def _text(value: str | None, label: str, limit: int, required: bool = False) -> str:
    text = (value or "").strip()
    if required and not text:
        raise AppError(400, f"{label} is required.")
    if len(text) > limit:
        raise AppError(400, f"{label} must be {limit} characters or fewer.")
    return text


def create_system(user: dict, body: dict) -> dict:
    _require_editor(user)
    code = _text(str(body.get("system_code") or ""), "System code", MAX_CODE, required=True).upper()
    if "," in code or " " in code:
        raise AppError(400, "System code cannot contain spaces or commas.")
    name = _text(str(body.get("system_name") or ""), "System name", MAX_NAME, required=True)
    description = _text(str(body.get("description") or ""), "Description", MAX_DESCRIPTION)
    known = existing_codes()
    if code in known:
        raise AppError(400, f"System code {code} already exists.")
    has_batch = _flag(body.get("has_batch"), False)
    has_report = _flag(body.get("has_report"), False)
    has_capability = _flag(body.get("has_capability"), False)
    now = _now()
    db.exec_sql(
        """
        INSERT INTO systems (
            system_code, system_name, description, upstreams, downstreams, has_batch, has_report, has_capability,
            created_at, last_modified
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [code, name, description, "", "", has_batch, has_report, has_capability, now, now],
    )
    _embed(_row(code))
    gcs.upload_db()
    return _payload(_row(code))


def update_system(user: dict, code: str, body: dict) -> dict:
    _require_editor(user)
    current = _row(code)
    code = current["system_code"]
    name = _text(str(body.get("system_name", current["system_name"]) or ""), "System name", MAX_NAME, required=True)
    description = _text(str(body.get("description", current["description"]) or ""), "Description", MAX_DESCRIPTION)
    has_batch = _flag(body["has_batch"], bool(current.get("has_batch"))) if "has_batch" in body else bool(current.get("has_batch"))
    has_report = _flag(body["has_report"], bool(current.get("has_report"))) if "has_report" in body else bool(current.get("has_report"))
    has_capability = (
        _flag(body["has_capability"], bool(current.get("has_capability")))
        if "has_capability" in body
        else bool(current.get("has_capability"))
    )
    db.exec_sql(
        """
        UPDATE systems
        SET system_name = ?, description = ?, upstreams = ?, downstreams = ?, has_batch = ?, has_report = ?,
            has_capability = ?, last_modified = ?
        WHERE system_code = ?
        """,
        [name, description, "", "", has_batch, has_report, has_capability, _now(), code],
    )
    _embed(_row(code))
    gcs.upload_db()
    return _payload(_row(code))


def _without(value: str | None, code: str) -> str:
    return ",".join(item for item in split_codes(value) if item != code)


def _impact_rows(code: str) -> dict:
    capabilities = db.fetchall(
        "SELECT capability_id, country, l1_capability, l2_capability, l3_capability FROM kb_capabilities WHERE system_code = ?"
        " ORDER BY l1_capability, l2_capability, l3_capability",
        [code],
    )
    batches = db.fetchall("SELECT batch_id, country, batch_name FROM kb_batches WHERE system_code = ? ORDER BY batch_name", [code])
    reports = db.fetchall("SELECT report_id, country, report_name FROM kb_reports WHERE system_code = ? ORDER BY report_name", [code])
    interfaces = [
        row
        for row in db.fetchall("SELECT interface_id, channel, country, interface_name, source, destination FROM interfaces ORDER BY interface_name")
        if code in split_codes(row["source"]) or code in split_codes(row["destination"])
    ]
    capability_ids = [row["capability_id"] for row in capabilities]
    capability_details = db.fetchall(
        "SELECT node_id, country, level, l1_capability, l2_capability FROM kb_capability_nodes WHERE system_code = ?"
        " ORDER BY level, l1_capability, l2_capability",
        [code],
    )
    unlinked_tables = []
    if capability_ids:
        marks = ", ".join("?" for _ in capability_ids)
        unlinked_tables = db.fetchall(
            f"""
            SELECT DISTINCT t.table_id, t.country, t.database_name, t.schema_name, t.table_name
            FROM kb_capability_tables l JOIN kb_tables t ON t.table_id = l.table_id
            WHERE l.capability_id IN ({marks})
            ORDER BY t.database_name, t.schema_name, t.table_name
            """,
            capability_ids,
        )
    peers = [
        row["system_code"]
        for row in db.fetchall("SELECT system_code, upstreams, downstreams FROM systems WHERE system_code <> ?", [code])
        if code in split_codes(row["upstreams"]) or code in split_codes(row["downstreams"])
    ]
    batch_peers = [
        row["batch_name"]
        for row in db.fetchall("SELECT batch_name, system_code, to_from_systems FROM kb_batches WHERE system_code <> ?", [code])
        if code in split_codes(row["to_from_systems"])
    ]
    projects = [
        row["project_name"]
        for row in db.fetchall("SELECT project_name, systems FROM projects ORDER BY project_name")
        if code in split_codes(row["systems"])
    ]
    return {
        "capabilities": capabilities,
        "capability_details": capability_details,
        "interfaces_kept": interfaces,
        "batches": batches,
        "reports": reports,
        "unlinked_tables": unlinked_tables,
        "systems_referencing": peers,
        "batches_referencing": batch_peers,
        "projects_in_scope": projects,
    }


def delete_impact(user: dict, code: str) -> dict:
    """Everything deleting the system removes or changes, for the warning shown before it happens."""
    require(user, "knowledge", "read")
    row = _row(code)
    impact = _impact_rows(row["system_code"])
    return {"system_code": row["system_code"], "system_name": row["system_name"], **impact}


def delete_system(user: dict, code: str, confirm: str = "") -> dict:
    """Delete the system and everything that belongs to it.

    Removed: its capabilities (and their links, and L1/L2 details), batches and reports. Interfaces that have it as source or destination
    are kept as they are, and tables are kept; both are only unlinked from the removed capabilities. The code is also
    taken out of other systems' upstreams / downstreams, other batches' To/From lists and project scopes.
    The caller must repeat the system code in `confirm`.
    """
    _require_editor(user)
    row = _row(code)
    code = row["system_code"]
    if confirm.strip().upper() != code:
        raise AppError(400, f"Type the system code {code} to confirm the delete")
    impact = _impact_rows(code)
    capability_ids = [item["capability_id"] for item in impact["capabilities"]]
    batch_ids = [item["batch_id"] for item in impact["batches"]]
    report_ids = [item["report_id"] for item in impact["reports"]]

    def delete_in(table: str, column: str, ids: list[str]) -> None:
        if ids:
            db.exec_sql(f"DELETE FROM {table} WHERE {column} IN ({', '.join('?' for _ in ids)})", ids)

    with db.transaction():
        for link in ("kb_capability_interfaces", "kb_capability_tables", "kb_capability_batches", "kb_capability_reports"):
            delete_in(link, "capability_id", capability_ids)
        delete_in("kb_capability_batches", "batch_id", batch_ids)
        delete_in("kb_capability_reports", "report_id", report_ids)
        delete_in("kb_capabilities", "capability_id", capability_ids)
        db.exec_sql("DELETE FROM kb_capability_nodes WHERE system_code = ?", [code])
        delete_in("kb_batches", "batch_id", batch_ids)
        delete_in("kb_reports", "report_id", report_ids)
        for other in db.fetchall("SELECT system_code, upstreams, downstreams FROM systems WHERE system_code <> ?", [code]):
            if code in split_codes(other["upstreams"]) or code in split_codes(other["downstreams"]):
                db.exec_sql(
                    "UPDATE systems SET upstreams = ?, downstreams = ? WHERE system_code = ?",
                    [_without(other["upstreams"], code), _without(other["downstreams"], code), other["system_code"]],
                )
        for batch in db.fetchall("SELECT batch_id, to_from_systems FROM kb_batches"):
            if code in split_codes(batch["to_from_systems"]):
                db.exec_sql("UPDATE kb_batches SET to_from_systems = ? WHERE batch_id = ?", [_without(batch["to_from_systems"], code), batch["batch_id"]])
        for project in db.fetchall("SELECT project_id, systems FROM projects"):
            if code in split_codes(project["systems"]):
                db.exec_sql("UPDATE projects SET systems = ? WHERE project_id = ?", [_without(project["systems"], code), project["project_id"]])
        db.exec_sql("DELETE FROM systems WHERE system_code = ?", [code])
    gcs.upload_db()
    return {
        "ok": True,
        "deleted": {
            "capabilities": len(capability_ids),
            "batches": len(batch_ids),
            "reports": len(report_ids),
        },
        "unlinked_tables": len(impact["unlinked_tables"]),
        "interfaces_kept": len(impact["interfaces_kept"]),
    }


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            if rows and any(cell.strip().upper() == "SYSTEM_CODE" for cell in rows[0]):
                return rows
        sheet = workbook.worksheets[0]
        return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppError(400, "Could not read the Excel file. Upload an .xlsx file.") from exc


def parse_systems_excel(data: bytes) -> list[dict]:
    """SYSTEM_CODE and SYSTEM_NAME are required. DESCRIPTION may be omitted or left blank."""
    rows = _read_sheet(data)
    if not rows:
        raise AppError(400, "The Excel file is empty.")
    header = [value.strip().upper() for value in rows[0]]
    missing = [name for name in ("SYSTEM_CODE", "SYSTEM_NAME") if name not in header]
    if missing:
        raise AppError(400, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(UPLOAD_HEADERS)}.")
    index = {name: header.index(name) if name in header else None for name in UPLOAD_HEADERS}
    records: list[dict] = []
    errors: list[dict] = []
    seen: dict[str, int] = {}

    def error(row_number: int, field: str, message: str) -> None:
        errors.append({"row": row_number, "field": field, "message": message})

    for offset, row in enumerate(rows[1:], start=2):
        values = {name: row[index[name]] if index[name] is not None and index[name] < len(row) else "" for name in UPLOAD_HEADERS}
        if not any(values.values()):
            continue
        code = values["SYSTEM_CODE"].strip().upper()
        name = values["SYSTEM_NAME"].strip()
        description = values["DESCRIPTION"].strip()
        if not code:
            error(offset, "SYSTEM_CODE", "SYSTEM_CODE is required.")
        elif "," in code or " " in code:
            error(offset, "SYSTEM_CODE", "SYSTEM_CODE cannot contain spaces or commas.")
        elif len(code) > MAX_CODE:
            error(offset, "SYSTEM_CODE", f"SYSTEM_CODE must be {MAX_CODE} characters or fewer.")
        if not name:
            error(offset, "SYSTEM_NAME", "SYSTEM_NAME is required.")
        elif len(name) > MAX_NAME:
            error(offset, "SYSTEM_NAME", f"SYSTEM_NAME must be {MAX_NAME} characters or fewer.")
        if len(description) > MAX_DESCRIPTION:
            error(offset, "DESCRIPTION", f"DESCRIPTION must be {MAX_DESCRIPTION} characters or fewer.")
        if code and code in seen:
            error(offset, "SYSTEM_CODE", f"Duplicate system code (first seen on row {seen[code]}).")
        elif code:
            seen[code] = offset
        records.append({"system_code": code, "system_name": name, "description": description})

    if len(records) > MAX_UPLOAD_ROWS:
        raise AppError(400, f"At most {MAX_UPLOAD_ROWS} rows can be uploaded at once.")
    if errors:
        raise AppError(400, f"{len(errors)} problem(s) found. Nothing was imported.", errors)
    if not records:
        raise AppError(400, "The Excel file has no data rows.")
    return records


def upload_from_excel(user: dict, filename: str, data: bytes) -> list[dict]:
    """Insert new systems and update the name of systems that already exist.

    A blank description is stored for a new system and left unchanged for one that already exists.
    The batch and report flags are not taken from the file.
    """
    _require_editor(user)
    if not (filename or "").lower().endswith(".xlsx"):
        raise AppError(400, "Upload an Excel file (.xlsx).")
    if len(data) > MAX_UPLOAD_BYTES:
        raise AppError(400, "The file is larger than 20 MB.")
    records = parse_systems_excel(data)
    known = existing_codes()
    now = _now()
    for rec in records:
        current = db.fetchone("SELECT * FROM systems WHERE system_code = ?", [rec["system_code"]]) if rec["system_code"] in known else None
        if current is None:
            db.exec_sql(
                """
                INSERT INTO systems (
                    system_code, system_name, description, upstreams, downstreams, has_batch, has_report, has_capability,
                    created_at, last_modified
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [rec["system_code"], rec["system_name"], rec["description"], "", "", False, False, False, now, now],
            )
        else:
            description = rec["description"] if rec["description"] else (current["description"] or "")
            db.exec_sql(
                """
                UPDATE systems
                SET system_name = ?, description = ?, last_modified = ?
                WHERE system_code = ?
                """,
                [rec["system_name"], description, now, rec["system_code"]],
            )
        _embed(_row(rec["system_code"]))
    gcs.upload_db()
    return list_systems()
