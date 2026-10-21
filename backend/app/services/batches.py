"""Structured knowledge-base batches organized by country and system."""

import io
import uuid
from datetime import datetime, timezone

from app import db
from app.config import kb_countries
from app.errors import AppError
from app.permissions import require
from app.services import gcs, llm, systems

HEADERS = ("BATCH_NAME", "BATCH_TYPE", "TO_FROM_SYSTEMS", "DESCRIPTION", "FREQUENCY", "BATCH_FILENAME", "FILE_LAYOUT")
BATCH_TYPES = ("incoming", "outgoing", "transaction")
MAX_ROWS = 5000
MAX_BYTES = 20 * 1024 * 1024
MAX_NAME = 200
MAX_DESCRIPTION = 4000
MAX_FILENAME = 500
MAX_LAYOUT = 50000
MAX_TO_FROM = 1000


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _require_editor(user: dict) -> None:
    require(user, "knowledge", "write")


def _canonical(value: str, allowed: tuple[str, ...]) -> str | None:
    lookup = {item.lower(): item for item in allowed}
    return lookup.get((value or "").strip().lower())


def _country(value: str) -> str:
    countries = kb_countries()
    canonical = _canonical(value, countries)
    if canonical is None:
        raise AppError(400, f"Country must be one of: {', '.join(countries)}.")
    return canonical


def _system_code(value: str, *, allow: str | None = None) -> str:
    code = (value or "").strip().upper()
    if not code:
        raise AppError(400, "System is required.")
    if code not in systems.existing_codes():
        raise AppError(400, f"System '{value}' is not a known system code.")
    allowed = (allow or "").strip().upper()
    if code != allowed and code not in systems.codes_with_batch():
        raise AppError(400, f"System '{code}' is not marked as having batches. Enable Has batch on the Systems page.")
    return code


FREQUENCIES = ("Adhoc", "Daily", "Weekly", "Monthly", "Quarterly", "Other")
MAX_FREQUENCY_OTHER = 100


def _frequency(body: dict, current: dict | None = None) -> tuple[str, str]:
    if "frequency" in body:
        raw = body.get("frequency")
    elif current is not None:
        raw = current.get("frequency") or ""
    else:
        raw = ""
    parts = [str(item).strip() for item in raw] if isinstance(raw, list) else [item.strip() for item in str(raw or "").split(",")]
    selected: list[str] = []
    for item in parts:
        if not item:
            continue
        canonical = next((name for name in FREQUENCIES if name.lower() == item.lower()), None)
        if canonical is None:
            raise AppError(400, "Frequency must be Adhoc, Daily, Weekly, Monthly, Quarterly, or Other.")
        if canonical not in selected:
            selected.append(canonical)
    if "frequency_other" in body:
        other = str(body.get("frequency_other") or "").strip()
    elif current is not None:
        other = str(current.get("frequency_other") or "").strip()
    else:
        other = ""
    if "Other" in selected:
        if not other:
            raise AppError(400, "Enter a frequency for Other.")
        if len(other) > MAX_FREQUENCY_OTHER:
            raise AppError(400, f"Other frequency must be {MAX_FREQUENCY_OTHER} characters or fewer.")
    else:
        other = ""
    ordered = [name for name in FREQUENCIES if name in selected]
    return ", ".join(ordered), other


def _frequency_cell(value: str) -> tuple[str, str]:
    """Excel FREQUENCY cell. Empty applies no frequency. Other is written as 'Other: note'."""
    selected: list[str] = []
    other = ""
    for item in str(value or "").split(","):
        text = item.strip()
        if not text:
            continue
        note = ""
        label = text
        if text.lower().startswith("other:"):
            label = "Other"
            note = text.split(":", 1)[1].strip()
        canonical = next((name for name in FREQUENCIES if name.lower() == label.lower()), None)
        if canonical is None:
            raise AppError(400, "Frequency must be Adhoc, Daily, Weekly, Monthly, Quarterly, or Other.")
        if canonical == "Other" and note:
            if other and other != note:
                raise AppError(400, "Enter one value for Other.")
            other = note
        if canonical not in selected:
            selected.append(canonical)
    if "Other" in selected:
        if not other:
            raise AppError(400, "Enter a frequency for Other, written as Other: note.")
        if len(other) > MAX_FREQUENCY_OTHER:
            raise AppError(400, f"Other frequency must be {MAX_FREQUENCY_OTHER} characters or fewer.")
    return ", ".join(name for name in FREQUENCIES if name in selected), other


def _batch_type(value: str) -> str:
    canonical = _canonical(value, BATCH_TYPES)
    if canonical is None:
        raise AppError(400, "Batch type must be incoming, outgoing, or transaction.")
    return canonical


def _to_from_systems(value: str | None) -> str:
    codes = systems.split_codes(value)
    if not codes:
        return ""
    known = systems.existing_codes()
    unknown = [code for code in codes if code not in known]
    if unknown:
        raise AppError(400, f"To/From system(s) has unknown system code(s): {', '.join(unknown)}.")
    joined = ", ".join(codes)
    if len(joined) > MAX_TO_FROM:
        raise AppError(400, f"To/From system(s) must be {MAX_TO_FROM} characters or fewer.")
    return joined


def text_for(row: dict) -> str:
    layout = (row.get("file_layout") or "")[:1500]
    desc = f" {row['description']}" if row.get("description") else ""
    filename = f" File: {row['batch_filename']}." if row.get("batch_filename") else ""
    peers = f" To/From: {row['to_from_systems']}." if row.get("to_from_systems") else ""
    return (
        f"Batch {row['batch_name']} ({row['batch_type']}, country {row['country']}, "
        f"system {row['system_code']}).{desc}{peers}{filename} "
        f"Frequency: {row.get('frequency') or 'unspecified'}"
        f"{' (' + row['frequency_other'] + ')' if row.get('frequency_other') else ''}. Layout: {layout}"
    )


def _embed(batch_id: str) -> None:
    row = db.fetchone("SELECT * FROM kb_batches WHERE batch_id = ?", [batch_id])
    if row is None:
        return
    vector = llm.embed(text_for(row)[:2000]) if llm.embedding_configured() else None
    db.exec_sql("UPDATE kb_batches SET embedding = ? WHERE batch_id = ?", [vector, batch_id])


def _payload(row: dict) -> dict:
    return {
        "batch_id": row["batch_id"],
        "country": row["country"],
        "system_code": row["system_code"],
        "batch_name": row["batch_name"],
        "batch_type": row["batch_type"],
        "frequency": row.get("frequency") or "",
        "frequency_other": row.get("frequency_other") or "",
        "to_from_systems": row.get("to_from_systems") or "",
        "description": row["description"],
        "batch_filename": row["batch_filename"],
        "file_layout": row["file_layout"],
        "created_at": row["created_at"],
        "last_modified": row["last_modified"],
    }


def list_batches() -> list[dict]:
    rows = db.fetchall(
        """
        SELECT * FROM kb_batches
        ORDER BY country, system_code, batch_name
        """
    )
    return [_payload(row) for row in rows]


def get_batch(batch_id: str) -> dict:
    row = db.fetchone("SELECT * FROM kb_batches WHERE batch_id = ?", [batch_id])
    if row is None:
        raise AppError(404, "Batch not found")
    return _payload(row)


def _validate_fields(
    name: str, batch_type: str, to_from: str, description: str, filename: str, layout: str
) -> tuple[str, str, str, str, str, str]:
    name = (name or "").strip()
    if not name:
        raise AppError(400, "Batch name is required.")
    if len(name) > MAX_NAME:
        raise AppError(400, f"Batch name must be {MAX_NAME} characters or fewer.")
    typed = _batch_type(batch_type)
    peers = _to_from_systems(to_from)
    description = (description or "").strip()
    filename = (filename or "").strip()
    layout = (layout or "").strip()
    if len(description) > MAX_DESCRIPTION:
        raise AppError(400, f"Description must be {MAX_DESCRIPTION} characters or fewer.")
    if len(filename) > MAX_FILENAME:
        raise AppError(400, f"Batch filename must be {MAX_FILENAME} characters or fewer.")
    if len(layout) > MAX_LAYOUT:
        raise AppError(400, f"File layout must be {MAX_LAYOUT} characters or fewer.")
    return name, typed, peers, description, filename, layout


def create_batch(user: dict, body: dict) -> dict:
    _require_editor(user)
    country = _country(str(body.get("country") or ""))
    system_code = _system_code(str(body.get("system_code") or body.get("system") or ""))
    name, batch_type, to_from, description, filename, layout = _validate_fields(
        str(body.get("batch_name") or ""),
        str(body.get("batch_type") or ""),
        str(body.get("to_from_systems") or ""),
        str(body.get("description") or ""),
        str(body.get("batch_filename") or ""),
        str(body.get("file_layout") or ""),
    )
    frequency, frequency_other = _frequency(body)
    existing = db.fetchone(
        """
        SELECT batch_id FROM kb_batches
        WHERE country = ? AND system_code = ? AND lower(batch_name) = lower(?)
        """,
        [country, system_code, name],
    )
    if existing is not None:
        raise AppError(400, f"Batch {name} already exists for {country} / {system_code}.")
    batch_id = str(uuid.uuid4())
    now = _now()
    db.exec_sql(
        """
        INSERT INTO kb_batches (
            batch_id, country, system_code, batch_name, batch_type, frequency, frequency_other, to_from_systems, description,
            batch_filename, file_layout, created_at, last_modified
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [batch_id, country, system_code, name, batch_type, frequency, frequency_other, to_from, description, filename, layout, now, now],
    )
    _embed(batch_id)
    gcs.upload_db()
    return get_batch(batch_id)


def update_batch(user: dict, batch_id: str, body: dict) -> dict:
    _require_editor(user)
    current = db.fetchone("SELECT * FROM kb_batches WHERE batch_id = ?", [batch_id])
    if current is None:
        raise AppError(404, "Batch not found")
    country = _country(str(body.get("country", current["country"]) or ""))
    system_code = _system_code(
        str(body.get("system_code", body.get("system", current["system_code"])) or ""),
        allow=current["system_code"],
    )
    name, batch_type, to_from, description, filename, layout = _validate_fields(
        str(body.get("batch_name", current["batch_name"]) or ""),
        str(body.get("batch_type", current["batch_type"]) or ""),
        str(body.get("to_from_systems", current.get("to_from_systems") or "") or ""),
        str(body.get("description", current["description"]) or ""),
        str(body.get("batch_filename", current["batch_filename"]) or ""),
        str(body.get("file_layout", current["file_layout"]) or ""),
    )
    frequency, frequency_other = _frequency(body, current)
    conflict = db.fetchone(
        """
        SELECT batch_id FROM kb_batches
        WHERE country = ? AND system_code = ? AND lower(batch_name) = lower(?) AND batch_id <> ?
        """,
        [country, system_code, name, batch_id],
    )
    if conflict is not None:
        raise AppError(400, f"Batch {name} already exists for {country} / {system_code}.")
    db.exec_sql(
        """
        UPDATE kb_batches
        SET country = ?, system_code = ?, batch_name = ?, batch_type = ?, frequency = ?, frequency_other = ?, to_from_systems = ?,
            description = ?, batch_filename = ?, file_layout = ?, last_modified = ?
        WHERE batch_id = ?
        """,
        [country, system_code, name, batch_type, frequency, frequency_other, to_from, description, filename, layout, _now(), batch_id],
    )
    _embed(batch_id)
    gcs.upload_db()
    return get_batch(batch_id)


def delete_batch(user: dict, batch_id: str) -> None:
    _require_editor(user)
    get_batch(batch_id)
    db.exec_sql("DELETE FROM kb_batches WHERE batch_id = ?", [batch_id])
    gcs.upload_db()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            if rows and any(h.strip().upper() == "BATCH_NAME" for h in rows[0]):
                return rows
        sheet = workbook.worksheets[0]
        return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppError(400, "Could not read the Excel file. Upload an .xlsx file.") from exc


def parse_and_validate(data: bytes) -> list[dict]:
    rows = _read_sheet(data)
    if not rows:
        raise AppError(400, "The Excel file is empty.")
    header = [value.strip().upper() for value in rows[0]]
    missing = [name for name in HEADERS if name not in header]
    if missing:
        raise AppError(400, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(HEADERS)}.")
    index = {name: header.index(name) for name in HEADERS}
    records: list[dict] = []
    errors: list[dict] = []
    seen: dict[str, int] = {}

    def error(row_number: int, field: str, message: str) -> None:
        errors.append({"row": row_number, "field": field, "message": message})

    known_codes = systems.existing_codes()
    for offset, row in enumerate(rows[1:], start=2):
        values = {name: row[index[name]] if index[name] < len(row) else "" for name in HEADERS}
        if not any(values.values()):
            continue
        name = values["BATCH_NAME"].strip()
        batch_type_raw = values["BATCH_TYPE"].strip()
        to_from_raw = values["TO_FROM_SYSTEMS"].strip()
        description = values["DESCRIPTION"].strip()
        filename = values["BATCH_FILENAME"].strip()
        layout = values["FILE_LAYOUT"].strip()
        frequency = ""
        frequency_other = ""
        try:
            frequency, frequency_other = _frequency_cell(values["FREQUENCY"])
        except AppError as exc:
            error(offset, "FREQUENCY", exc.detail)
        if not name:
            error(offset, "BATCH_NAME", "BATCH_NAME is required.")
        key = name.upper()
        if name and key in seen:
            error(offset, "BATCH_NAME", f"Duplicate batch name '{name}' (first seen on row {seen[key]}).")
        elif name:
            seen[key] = offset
        batch_type = _canonical(batch_type_raw, BATCH_TYPES)
        if not batch_type_raw:
            error(offset, "BATCH_TYPE", "BATCH_TYPE is required.")
        elif batch_type is None:
            error(offset, "BATCH_TYPE", "BATCH_TYPE must be incoming, outgoing, or transaction.")
        peer_codes = systems.split_codes(to_from_raw)
        unknown = [code for code in peer_codes if code not in known_codes]
        if unknown:
            error(offset, "TO_FROM_SYSTEMS", f"Unknown system code(s): {', '.join(unknown)}.")
        records.append(
            {
                "batch_name": name,
                "batch_type": batch_type or batch_type_raw.lower(),
                "to_from_systems": ", ".join(peer_codes),
                "description": description,
                "frequency": frequency,
                "frequency_other": frequency_other,
                "batch_filename": filename,
                "file_layout": layout,
            }
        )

    if len(records) > MAX_ROWS:
        raise AppError(400, f"At most {MAX_ROWS} rows can be uploaded at once.")
    if errors:
        raise AppError(400, f"{len(errors)} problem(s) found. Nothing was imported.", errors)
    if not records:
        raise AppError(400, "The Excel file has no data rows.")
    return records


def upload_from_excel(user: dict, country_raw: str, system_raw: str, filename: str, data: bytes) -> dict:
    _require_editor(user)
    country = _country(country_raw)
    system_code = _system_code(system_raw)
    if not (filename or "").lower().endswith(".xlsx"):
        raise AppError(400, "Upload an Excel file (.xlsx).")
    if len(data) > MAX_BYTES:
        raise AppError(400, "The file is larger than 20 MB.")
    records = parse_and_validate(data)
    now = _now()
    for rec in records:
        existing = db.fetchone(
            """
            SELECT batch_id, description, batch_filename, file_layout, batch_type, to_from_systems, frequency, frequency_other
            FROM kb_batches
            WHERE country = ? AND system_code = ? AND lower(batch_name) = lower(?)
            """,
            [country, system_code, rec["batch_name"]],
        )
        if existing is None:
            batch_id = str(uuid.uuid4())
            db.exec_sql(
                """
                INSERT INTO kb_batches (
                    batch_id, country, system_code, batch_name, batch_type, frequency, frequency_other, to_from_systems, description,
                    batch_filename, file_layout, created_at, last_modified
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    batch_id,
                    country,
                    system_code,
                    rec["batch_name"],
                    rec["batch_type"],
                    rec["frequency"],
                    rec["frequency_other"],
                    rec["to_from_systems"],
                    rec["description"],
                    rec["batch_filename"],
                    rec["file_layout"],
                    now,
                    now,
                ],
            )
            _embed(batch_id)
        else:
            description = rec["description"] or existing["description"]
            batch_filename = rec["batch_filename"] or existing["batch_filename"]
            file_layout = rec["file_layout"] or existing["file_layout"]
            to_from = rec["to_from_systems"] or existing.get("to_from_systems") or ""
            frequency = rec["frequency"] or existing["frequency"] or ""
            frequency_other = rec["frequency_other"] if rec["frequency"] else (existing["frequency_other"] or "")
            db.exec_sql(
                """
                UPDATE kb_batches
                SET batch_name = ?, batch_type = ?, frequency = ?, frequency_other = ?, to_from_systems = ?, description = ?,
                    batch_filename = ?, file_layout = ?, last_modified = ?
                WHERE batch_id = ?
                """,
                [
                    rec["batch_name"],
                    rec["batch_type"],
                    frequency,
                    frequency_other,
                    to_from,
                    description,
                    batch_filename,
                    file_layout,
                    now,
                    existing["batch_id"],
                ],
            )
            _embed(existing["batch_id"])

    from app.services import reindex

    reindex.embed_missing_in_background()
    gcs.upload_db()
    return {"batches": list_batches(), "uploaded": len(records)}
