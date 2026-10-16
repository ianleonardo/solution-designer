"""Knowledge base database tables and columns, organized by country, database and schema."""

import io
import json
import re
import uuid
from datetime import datetime, timezone

from app import db
from app.config import KB_EDITORS
from app.errors import AppError
from app.services import gcs, llm

COUNTRIES = ("Singapore", "Malaysia", "Thailand", "Indonesia", "Vietnam")
DATABASES = ("PIB", "MBK", "SMS", "NTN", "LCM")
HEADERS = ("TABLE_NAME", "COLUMN_NAME", "DATA_TYPE", "IS_NULL", "PRIMARY_KEY", "FOREIGN_KEY")
MAX_ROWS = 5000
MAX_BYTES = 20 * 1024 * 1024
MAX_NAME = 200
MAX_DESCRIPTION = 2000


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _require_editor(user: dict) -> None:
    if user["role_name"] not in KB_EDITORS:
        raise AppError(403, "Forbidden")


def _canonical(value: str, allowed: tuple[str, ...]) -> str | None:
    lookup = {item.lower(): item for item in allowed}
    return lookup.get((value or "").strip().lower())


def _country(value: str) -> str:
    canonical = _canonical(value, COUNTRIES)
    if canonical is None:
        raise AppError(400, f"Country must be one of: {', '.join(COUNTRIES)}.")
    return canonical


def _database(value: str) -> str:
    canonical = _canonical(value, DATABASES)
    if canonical is None:
        raise AppError(400, f"Database must be one of: {', '.join(DATABASES)}.")
    return canonical


def _schema(value: str | None, database: str) -> str:
    cleaned = (value or "").strip()
    if not cleaned:
        return database
    if len(cleaned) > MAX_NAME:
        raise AppError(400, f"Schema must be {MAX_NAME} characters or fewer.")
    return cleaned


def _clean_text(value: str | None, label: str, limit: int = MAX_NAME, required: bool = False) -> str:
    cleaned = (value or "").strip()
    if required and not cleaned:
        raise AppError(400, f"{label} is required.")
    if len(cleaned) > limit:
        raise AppError(400, f"{label} must be {limit} characters or fewer.")
    return cleaned


def _clean_is_null(value: str | None) -> str:
    cleaned = (value or "").strip().upper()
    if not cleaned:
        return "N"
    if cleaned in {"Y", "YES", "TRUE", "1", "NULLABLE"}:
        return "Y"
    if cleaned in {"N", "NO", "FALSE", "0", "NOT NULL"}:
        return "N"
    return cleaned[:10]


def _human_identifier(name: str) -> str:
    words = re.sub(r"[_-]+", " ", name.strip())
    words = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", words)
    return " ".join(words.split()).lower() or name


def heuristic_table_description(table_name: str, columns: list[dict]) -> str:
    human = _human_identifier(table_name)
    col_names = [col["column_name"] for col in columns[:5] if col.get("column_name")]
    suffix = f" (key attributes: {', '.join(col_names)})" if col_names else ""
    return f"Stores {human} records{suffix}."


def heuristic_column_description(table_name: str, column: dict) -> str:
    col_name = column["column_name"]
    human = _human_identifier(col_name)
    data_type = column.get("data_type") or ""
    pk = (column.get("primary_key") or "").strip()
    fk = (column.get("foreign_key") or "").strip()
    notes = []
    if pk and pk.upper() not in {"N", "NO", "FALSE", "0"}:
        notes.append("primary key")
    if fk and fk.upper() not in {"N", "NO", "FALSE", "0"}:
        notes.append(f"foreign key ({fk})" if fk.upper() not in {"Y", "YES", "TRUE", "1", "FK"} else "foreign key")
    if column.get("is_null", "N") == "N":
        notes.append("required")
    meta = f" [{', '.join(notes)}]" if notes else ""
    return f"{human[:1].upper() + human[1:]} of {_human_identifier(table_name)} ({data_type}){meta}."


def infer_descriptions(tables_spec: dict[str, list[dict]]) -> tuple[dict[str, str], dict[tuple[str, str], str]]:
    """Return ({table_name_upper: description}, {(table_name_upper, col_name_upper): description})."""
    table_descs: dict[str, str] = {}
    col_descs: dict[tuple[str, str], str] = {}

    for t_name, cols in tables_spec.items():
        t_key = t_name.upper()
        table_descs[t_key] = heuristic_table_description(t_name, cols)
        for col in cols:
            col_descs[(t_key, col["column_name"].upper())] = heuristic_column_description(t_name, col)

    if not llm.configured() or not tables_spec:
        return table_descs, col_descs

    try:
        prompt_items = []
        for t_name, cols in list(tables_spec.items())[:25]:
            col_list = [
                {
                    "column_name": c["column_name"],
                    "data_type": c["data_type"],
                    "is_null": c["is_null"],
                    "primary_key": c.get("primary_key", ""),
                    "foreign_key": c.get("foreign_key", ""),
                }
                for c in cols[:60]
            ]
            prompt_items.append({"table_name": t_name, "columns": col_list})
        raw = llm.complete(
            "You are a database architect. Infer concise, factual table purpose descriptions and column descriptions "
            "from table and column metadata. Return JSON only: "
            '{"tables":[{"table_name":"...","description":"...","columns":[{"column_name":"...","description":"..."}]}]}.',
            f"TABLES:\n{json.dumps(prompt_items)}",
            0.1,
        )
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end != -1:
            parsed = json.loads(raw[start : end + 1])
            for item in parsed.get("tables") or []:
                t_key = str(item.get("table_name") or "").strip().upper()
                t_desc = str(item.get("description") or "").strip()
                if t_key and t_desc:
                    table_descs[t_key] = t_desc[:MAX_DESCRIPTION]
                for c_item in item.get("columns") or []:
                    c_key = str(c_item.get("column_name") or "").strip().upper()
                    c_desc = str(c_item.get("description") or "").strip()
                    if t_key and c_key and c_desc:
                        col_descs[(t_key, c_key)] = c_desc[:MAX_DESCRIPTION]
    except Exception:
        pass

    return table_descs, col_descs


def text_for(row: dict, columns: list[dict] | None = None) -> str:
    """Readable summary of a table and its columns for embedding, search and chat context."""
    if columns is None:
        columns = db.fetchall(
            "SELECT * FROM kb_table_columns WHERE table_id = ? ORDER BY created_at, column_name",
            [row["table_id"]],
        )
    header = (
        f"Table {row['table_name']} (country {row['country']}, database {row['database_name']}, "
        f"schema {row['schema_name']})."
    )
    if row.get("description"):
        header += f" {row['description']}"
    if not columns:
        return header
    col_parts = []
    for col in columns[:60]:
        flags = [col["data_type"]]
        if col.get("is_null") == "N":
            flags.append("NOT NULL")
        if col.get("primary_key"):
            flags.append(f"PK:{col['primary_key']}")
        if col.get("foreign_key"):
            flags.append(f"FK:{col['foreign_key']}")
        desc = f" - {col['description']}" if col.get("description") else ""
        col_parts.append(f"{col['column_name']} ({', '.join(flags)}){desc}")
    return f"{header} Columns: {'; '.join(col_parts)}."


def _embed_table(table_id: str) -> None:
    row = db.fetchone("SELECT * FROM kb_tables WHERE table_id = ?", [table_id])
    if row is None:
        return
    cols = db.fetchall("SELECT * FROM kb_table_columns WHERE table_id = ? ORDER BY created_at, column_name", [table_id])
    vector = llm.embed(text_for(row, cols)[:2000]) if llm.embedding_configured() else None
    db.exec_sql("UPDATE kb_tables SET embedding = ? WHERE table_id = ?", [vector, table_id])


def _column_payload(row: dict) -> dict:
    return {
        "column_id": row["column_id"],
        "table_id": row["table_id"],
        "column_name": row["column_name"],
        "data_type": row["data_type"],
        "is_null": row["is_null"],
        "primary_key": row["primary_key"],
        "foreign_key": row["foreign_key"],
        "description": row["description"],
        "created_at": row["created_at"],
        "last_modified": row["last_modified"],
    }


def _table_payload(row: dict, columns: list[dict] | None = None, column_count: int | None = None) -> dict:
    payload = {
        "table_id": row["table_id"],
        "country": row["country"],
        "database_name": row["database_name"],
        "schema_name": row["schema_name"],
        "table_name": row["table_name"],
        "description": row["description"],
        "column_count": column_count if column_count is not None else (len(columns) if columns is not None else 0),
        "created_at": row["created_at"],
        "last_modified": row["last_modified"],
    }
    if columns is not None:
        payload["columns"] = [_column_payload(col) for col in columns]
    return payload


def list_tables() -> list[dict]:
    rows = db.fetchall(
        """
        SELECT t.*, (SELECT COUNT(*) FROM kb_table_columns c WHERE c.table_id = t.table_id) AS column_count
        FROM kb_tables t
        ORDER BY t.country, t.database_name, t.schema_name, t.table_name
        """
    )
    return [_table_payload(row, column_count=int(row.get("column_count") or 0)) for row in rows]


def _require_table(table_id: str) -> dict:
    row = db.fetchone("SELECT * FROM kb_tables WHERE table_id = ?", [table_id])
    if row is None:
        raise AppError(404, "Table not found")
    return row


def get_table(table_id: str) -> dict:
    row = _require_table(table_id)
    cols = db.fetchall(
        "SELECT * FROM kb_table_columns WHERE table_id = ? ORDER BY created_at, column_name",
        [table_id],
    )
    return _table_payload(row, columns=cols)


def create_table(user: dict, body: dict) -> dict:
    _require_editor(user)
    country = _country(str(body.get("country") or ""))
    database = _database(str(body.get("database_name") or body.get("database") or ""))
    schema = _schema(body.get("schema_name") if "schema_name" in body else body.get("schema"), database)
    table_name = _clean_text(str(body.get("table_name") or ""), "Table name", MAX_NAME, required=True)
    description = _clean_text(str(body.get("description") or ""), "Description", MAX_DESCRIPTION)

    existing = db.fetchone(
        """
        SELECT table_id FROM kb_tables
        WHERE country = ? AND database_name = ? AND lower(schema_name) = lower(?) AND lower(table_name) = lower(?)
        """,
        [country, database, schema, table_name],
    )
    if existing is not None:
        raise AppError(400, f"Table {table_name} already exists in {country} / {database} / {schema}.")

    if not description:
        description = heuristic_table_description(table_name, [])

    table_id = str(uuid.uuid4())
    now = _now()
    db.exec_sql(
        """
        INSERT INTO kb_tables (
            table_id, country, database_name, schema_name, table_name, description, created_at, last_modified
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [table_id, country, database, schema, table_name, description, now, now],
    )
    _embed_table(table_id)
    gcs.upload_db()
    return get_table(table_id)


def update_table(user: dict, table_id: str, body: dict) -> dict:
    _require_editor(user)
    current = _require_table(table_id)
    table_name = _clean_text(
        str(body.get("table_name", current["table_name"]) or ""),
        "Table name",
        MAX_NAME,
        required=True,
    )
    description = _clean_text(
        str(body.get("description", current["description"]) or ""),
        "Description",
        MAX_DESCRIPTION,
    )
    conflict = db.fetchone(
        """
        SELECT table_id FROM kb_tables
        WHERE country = ? AND database_name = ? AND lower(schema_name) = lower(?)
          AND lower(table_name) = lower(?) AND table_id <> ?
        """,
        [current["country"], current["database_name"], current["schema_name"], table_name, table_id],
    )
    if conflict is not None:
        raise AppError(
            400,
            f"Table {table_name} already exists in {current['country']} / {current['database_name']} / {current['schema_name']}.",
        )
    db.exec_sql(
        "UPDATE kb_tables SET table_name = ?, description = ?, last_modified = ? WHERE table_id = ?",
        [table_name, description, _now(), table_id],
    )
    _embed_table(table_id)
    gcs.upload_db()
    return get_table(table_id)


def delete_table(user: dict, table_id: str) -> None:
    _require_editor(user)
    _require_table(table_id)
    db.exec_sql("DELETE FROM kb_table_columns WHERE table_id = ?", [table_id])
    db.exec_sql("DELETE FROM kb_tables WHERE table_id = ?", [table_id])
    gcs.upload_db()


def add_column(user: dict, table_id: str, body: dict) -> dict:
    _require_editor(user)
    table = _require_table(table_id)
    column_name = _clean_text(str(body.get("column_name") or ""), "Column name", MAX_NAME, required=True)
    data_type = _clean_text(str(body.get("data_type") or ""), "Data type", MAX_NAME, required=True)
    is_null = _clean_is_null(body.get("is_null"))
    primary_key = _clean_text(str(body.get("primary_key") or ""), "Primary key", MAX_NAME)
    foreign_key = _clean_text(str(body.get("foreign_key") or ""), "Foreign key", MAX_NAME)
    description = _clean_text(str(body.get("description") or ""), "Description", MAX_DESCRIPTION)

    conflict = db.fetchone(
        "SELECT column_id FROM kb_table_columns WHERE table_id = ? AND lower(column_name) = lower(?)",
        [table_id, column_name],
    )
    if conflict is not None:
        raise AppError(400, f"Column {column_name} already exists on table {table['table_name']}.")

    if not description:
        description = heuristic_column_description(
            table["table_name"],
            {
                "column_name": column_name,
                "data_type": data_type,
                "is_null": is_null,
                "primary_key": primary_key,
                "foreign_key": foreign_key,
            },
        )

    column_id = str(uuid.uuid4())
    now = _now()
    db.exec_sql(
        """
        INSERT INTO kb_table_columns (
            column_id, table_id, column_name, data_type, is_null, primary_key, foreign_key, description, created_at, last_modified
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [column_id, table_id, column_name, data_type, is_null, primary_key, foreign_key, description, now, now],
    )
    db.exec_sql("UPDATE kb_tables SET last_modified = ? WHERE table_id = ?", [now, table_id])
    _embed_table(table_id)
    gcs.upload_db()
    row = db.fetchone("SELECT * FROM kb_table_columns WHERE column_id = ?", [column_id])
    assert row is not None
    return _column_payload(row)


def update_column(user: dict, table_id: str, column_id: str, body: dict) -> dict:
    _require_editor(user)
    table = _require_table(table_id)
    current = db.fetchone(
        "SELECT * FROM kb_table_columns WHERE column_id = ? AND table_id = ?",
        [column_id, table_id],
    )
    if current is None:
        raise AppError(404, "Column not found")
    column_name = _clean_text(
        str(body.get("column_name", current["column_name"]) or ""),
        "Column name",
        MAX_NAME,
        required=True,
    )
    data_type = _clean_text(
        str(body.get("data_type", current["data_type"]) or ""),
        "Data type",
        MAX_NAME,
        required=True,
    )
    is_null = _clean_is_null(body.get("is_null", current["is_null"]))
    primary_key = _clean_text(str(body.get("primary_key", current["primary_key"]) or ""), "Primary key", MAX_NAME)
    foreign_key = _clean_text(str(body.get("foreign_key", current["foreign_key"]) or ""), "Foreign key", MAX_NAME)
    description = _clean_text(str(body.get("description", current["description"]) or ""), "Description", MAX_DESCRIPTION)

    conflict = db.fetchone(
        "SELECT column_id FROM kb_table_columns WHERE table_id = ? AND lower(column_name) = lower(?) AND column_id <> ?",
        [table_id, column_name, column_id],
    )
    if conflict is not None:
        raise AppError(400, f"Column {column_name} already exists on table {table['table_name']}.")

    now = _now()
    db.exec_sql(
        """
        UPDATE kb_table_columns
        SET column_name = ?, data_type = ?, is_null = ?, primary_key = ?, foreign_key = ?, description = ?, last_modified = ?
        WHERE column_id = ?
        """,
        [column_name, data_type, is_null, primary_key, foreign_key, description, now, column_id],
    )
    db.exec_sql("UPDATE kb_tables SET last_modified = ? WHERE table_id = ?", [now, table_id])
    _embed_table(table_id)
    gcs.upload_db()
    row = db.fetchone("SELECT * FROM kb_table_columns WHERE column_id = ?", [column_id])
    assert row is not None
    return _column_payload(row)


def delete_column(user: dict, table_id: str, column_id: str) -> None:
    _require_editor(user)
    _require_table(table_id)
    current = db.fetchone(
        "SELECT column_id FROM kb_table_columns WHERE column_id = ? AND table_id = ?",
        [column_id, table_id],
    )
    if current is None:
        raise AppError(404, "Column not found")
    db.exec_sql("DELETE FROM kb_table_columns WHERE column_id = ?", [column_id])
    db.exec_sql("UPDATE kb_tables SET last_modified = ? WHERE table_id = ?", [_now(), table_id])
    _embed_table(table_id)
    gcs.upload_db()


def delete_columns_batch(user: dict, column_ids: list[str]) -> dict:
    _require_editor(user)
    cleaned = [str(cid).strip() for cid in column_ids if str(cid).strip()]
    if not cleaned:
        return {"deleted": 0}
    affected_tables: set[str] = set()
    deleted = 0
    for cid in cleaned:
        row = db.fetchone("SELECT table_id FROM kb_table_columns WHERE column_id = ?", [cid])
        if row is not None:
            affected_tables.add(row["table_id"])
            db.exec_sql("DELETE FROM kb_table_columns WHERE column_id = ?", [cid])
            deleted += 1
    now = _now()
    for table_id in affected_tables:
        db.exec_sql("UPDATE kb_tables SET last_modified = ? WHERE table_id = ?", [now, table_id])
        _embed_table(table_id)
    gcs.upload_db()
    return {"deleted": deleted}


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            if rows and any(h.strip().upper() == "TABLE_NAME" for h in rows[0]):
                return rows
        sheet = workbook.worksheets[0]
        return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppError(400, "Could not read the Excel file. Upload an .xlsx file.") from exc


def parse_and_validate(data: bytes) -> list[dict]:
    rows = _read_sheet(data)
    if not rows:
        raise AppError(400, "The Excel file is empty.")
    header = [value.strip().upper() for value in rows[0]]
    missing = [name for name in HEADERS if name not in header]
    if missing:
        raise AppError(400, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(HEADERS)}.")
    index = {name: header.index(name) for name in HEADERS}
    records: list[dict] = []
    errors: list[dict] = []
    seen_pairs: dict[tuple[str, str], int] = {}

    def error(row_number: int, field: str, message: str) -> None:
        errors.append({"row": row_number, "field": field, "message": message})

    for offset, row in enumerate(rows[1:], start=2):
        values = {name: row[index[name]] if index[name] < len(row) else "" for name in HEADERS}
        if not any(values.values()):
            continue
        table_name = values["TABLE_NAME"].strip()
        column_name = values["COLUMN_NAME"].strip()
        data_type = values["DATA_TYPE"].strip()
        if not table_name:
            error(offset, "TABLE_NAME", "TABLE_NAME is required.")
        if not column_name:
            error(offset, "COLUMN_NAME", "COLUMN_NAME is required.")
        if not data_type:
            error(offset, "DATA_TYPE", "DATA_TYPE is required.")
        if table_name and column_name:
            pair = (table_name.upper(), column_name.upper())
            if pair in seen_pairs:
                error(
                    offset,
                    "COLUMN_NAME",
                    f"Duplicate column '{column_name}' for table '{table_name}' (first seen on row {seen_pairs[pair]}).",
                )
            else:
                seen_pairs[pair] = offset
        records.append(
            {
                "table_name": table_name,
                "column_name": column_name,
                "data_type": data_type,
                "is_null": _clean_is_null(values["IS_NULL"]),
                "primary_key": values["PRIMARY_KEY"].strip(),
                "foreign_key": values["FOREIGN_KEY"].strip(),
            }
        )

    if len(records) > MAX_ROWS:
        raise AppError(400, f"At most {MAX_ROWS} rows can be uploaded at once.")
    if errors:
        raise AppError(400, f"{len(errors)} problem(s) found. Nothing was imported.", errors)
    if not records:
        raise AppError(400, "The Excel file has no data rows.")
    return records


def upload_from_excel(
    user: dict,
    country_raw: str,
    database_raw: str,
    schema_raw: str | None,
    filename: str,
    data: bytes,
) -> dict:
    _require_editor(user)
    country = _country(country_raw)
    database = _database(database_raw)
    schema = _schema(schema_raw, database)
    if not (filename or "").lower().endswith(".xlsx"):
        raise AppError(400, "Upload an Excel file (.xlsx).")
    if len(data) > MAX_BYTES:
        raise AppError(400, "The file is larger than 20 MB.")

    records = parse_and_validate(data)
    grouped: dict[str, list[dict]] = {}
    canonical_table_name: dict[str, str] = {}
    for rec in records:
        t_key = rec["table_name"].upper()
        canonical_table_name.setdefault(t_key, rec["table_name"])
        grouped.setdefault(canonical_table_name[t_key], []).append(rec)

    table_descs, col_descs = infer_descriptions(grouped)
    now = _now()
    missing_columns: list[dict] = []
    touched_table_ids: list[str] = []

    for t_name, cols in grouped.items():
        t_key = t_name.upper()
        inferred_t_desc = table_descs.get(t_key, heuristic_table_description(t_name, cols))
        existing_table = db.fetchone(
            """
            SELECT * FROM kb_tables
            WHERE country = ? AND database_name = ? AND lower(schema_name) = lower(?) AND lower(table_name) = lower(?)
            """,
            [country, database, schema, t_name],
        )
        if existing_table is None:
            table_id = str(uuid.uuid4())
            db.exec_sql(
                """
                INSERT INTO kb_tables (
                    table_id, country, database_name, schema_name, table_name, description, created_at, last_modified
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [table_id, country, database, schema, t_name, inferred_t_desc, now, now],
            )
        else:
            table_id = existing_table["table_id"]
            new_desc = existing_table["description"] or inferred_t_desc
            db.exec_sql(
                "UPDATE kb_tables SET description = ?, last_modified = ? WHERE table_id = ?",
                [new_desc, now, table_id],
            )
        touched_table_ids.append(table_id)

        existing_cols = db.fetchall("SELECT * FROM kb_table_columns WHERE table_id = ?", [table_id])
        existing_by_upper = {row["column_name"].upper(): row for row in existing_cols}
        uploaded_uppers: set[str] = set()

        for col in cols:
            c_key = col["column_name"].upper()
            uploaded_uppers.add(c_key)
            inferred_c_desc = col_descs.get((t_key, c_key), heuristic_column_description(t_name, col))
            current_col = existing_by_upper.get(c_key)
            if current_col is None:
                db.exec_sql(
                    """
                    INSERT INTO kb_table_columns (
                        column_id, table_id, column_name, data_type, is_null, primary_key, foreign_key,
                        description, created_at, last_modified
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        str(uuid.uuid4()),
                        table_id,
                        col["column_name"],
                        col["data_type"],
                        col["is_null"],
                        col["primary_key"],
                        col["foreign_key"],
                        inferred_c_desc,
                        now,
                        now,
                    ],
                )
            else:
                kept_desc = current_col["description"] or inferred_c_desc
                db.exec_sql(
                    """
                    UPDATE kb_table_columns
                    SET column_name = ?, data_type = ?, is_null = ?, primary_key = ?, foreign_key = ?,
                        description = ?, last_modified = ?
                    WHERE column_id = ?
                    """,
                    [
                        col["column_name"],
                        col["data_type"],
                        col["is_null"],
                        col["primary_key"],
                        col["foreign_key"],
                        kept_desc,
                        now,
                        current_col["column_id"],
                    ],
                )

        for c_upper, db_col in existing_by_upper.items():
            if c_upper not in uploaded_uppers:
                missing_columns.append(
                    {
                        "column_id": db_col["column_id"],
                        "table_id": table_id,
                        "table_name": existing_table["table_name"] if existing_table else t_name,
                        "column_name": db_col["column_name"],
                        "data_type": db_col["data_type"],
                        "description": db_col["description"],
                    }
                )

    from app.services import reindex

    reindex.embed_missing_in_background()
    for tid in touched_table_ids:
        _embed_table(tid)
    gcs.upload_db()

    return {
        "tables": list_tables(),
        "uploaded_tables": len(grouped),
        "uploaded_columns": len(records),
        "missing_columns": missing_columns,
    }
