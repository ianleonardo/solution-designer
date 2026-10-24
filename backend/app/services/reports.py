"""Structured knowledge-base reports organized by country and system."""

import io
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from app import db
from app.config import data_dir, kb_countries
from app.errors import AppError
from app.permissions import require
from app.services import gcs, llm, systems

HEADERS = ("REPORT_ID", "REPORT_NAME", "REPORT_DESCRIPTION", "FREQUENCY", "REPORT_LAYOUT")
MAX_ROWS = 5000
MAX_BYTES = 20 * 1024 * 1024
MAX_CODE = 50
MAX_NAME = 200
MAX_DESCRIPTION = 4000
MAX_LAYOUT = 50000
MAX_SAMPLE_BYTES = 20 * 1024 * 1024
SAMPLE_TYPES = {"png": "image/png", "jpg": "image/jpeg", "gif": "image/gif", "webp": "image/webp", "pdf": "application/pdf", "txt": "text/plain"}
_STORED_NAME = re.compile(r"^[A-Za-z0-9._-]{1,200}$")


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _require_editor(user: dict) -> None:
    require(user, "knowledge", "write")


def _canonical(value: str, allowed: tuple[str, ...]) -> str | None:
    lookup = {item.lower(): item for item in allowed}
    return lookup.get((value or "").strip().lower())


def _country(value: str) -> str:
    countries = kb_countries()
    canonical = _canonical(value, countries)
    if canonical is None:
        raise AppError(400, f"Country must be one of: {', '.join(countries)}.")
    return canonical


def _system_code(value: str, *, allow: str | None = None) -> str:
    code = (value or "").strip().upper()
    if not code:
        raise AppError(400, "System is required.")
    if code not in systems.existing_codes():
        raise AppError(400, f"System '{value}' is not a known system code.")
    allowed = (allow or "").strip().upper()
    if code != allowed and code not in systems.codes_with_report():
        raise AppError(400, f"System '{code}' is not marked as having reports. Enable Has report on the Systems page.")
    return code


FREQUENCIES = ("Adhoc", "Daily", "Weekly", "Monthly", "Quarterly", "Other")
MAX_FREQUENCY_OTHER = 100


def _frequency(body: dict, current: dict | None = None) -> tuple[str, str]:
    if "frequency" in body:
        raw = body.get("frequency")
    elif current is not None:
        raw = current.get("frequency") or ""
    else:
        raw = ""
    parts = [str(item).strip() for item in raw] if isinstance(raw, list) else [item.strip() for item in str(raw or "").split(",")]
    selected: list[str] = []
    for item in parts:
        if not item:
            continue
        canonical = next((name for name in FREQUENCIES if name.lower() == item.lower()), None)
        if canonical is None:
            raise AppError(400, "Frequency must be Adhoc, Daily, Weekly, Monthly, Quarterly, or Other.")
        if canonical not in selected:
            selected.append(canonical)
    if "frequency_other" in body:
        other = str(body.get("frequency_other") or "").strip()
    elif current is not None:
        other = str(current.get("frequency_other") or "").strip()
    else:
        other = ""
    if "Other" in selected:
        if not other:
            raise AppError(400, "Enter a frequency for Other.")
        if len(other) > MAX_FREQUENCY_OTHER:
            raise AppError(400, f"Other frequency must be {MAX_FREQUENCY_OTHER} characters or fewer.")
    else:
        other = ""
    ordered = [name for name in FREQUENCIES if name in selected]
    return ", ".join(ordered), other


def _frequency_cell(value: str) -> tuple[str, str]:
    """Excel FREQUENCY cell. Empty applies no frequency. Other is written as 'Other: note'."""
    selected: list[str] = []
    other = ""
    for item in str(value or "").split(","):
        text = item.strip()
        if not text:
            continue
        note = ""
        label = text
        if text.lower().startswith("other:"):
            label = "Other"
            note = text.split(":", 1)[1].strip()
        canonical = next((name for name in FREQUENCIES if name.lower() == label.lower()), None)
        if canonical is None:
            raise AppError(400, "Frequency must be Adhoc, Daily, Weekly, Monthly, Quarterly, or Other.")
        if canonical == "Other" and note:
            if other and other != note:
                raise AppError(400, "Enter one value for Other.")
            other = note
        if canonical not in selected:
            selected.append(canonical)
    if "Other" in selected:
        if not other:
            raise AppError(400, "Enter a frequency for Other, written as Other: note.")
        if len(other) > MAX_FREQUENCY_OTHER:
            raise AppError(400, f"Other frequency must be {MAX_FREQUENCY_OTHER} characters or fewer.")
    return ", ".join(name for name in FREQUENCIES if name in selected), other


def _report_code(value: str) -> str:
    code = (value or "").strip()
    if not code:
        raise AppError(400, "Report ID is required.")
    if len(code) > MAX_CODE:
        raise AppError(400, f"Report ID must be {MAX_CODE} characters or fewer.")
    return code


def text_for(row: dict) -> str:
    layout = (row.get("layout") or "")[:1500]
    desc = f" {row['description']}" if row.get("description") else ""
    code = f"{row['report_code']} " if row.get("report_code") else ""
    sample = _sample_public(row.get("sample") or "")
    sample_note = f" Sample: {sample['filename']}." if sample else ""
    return (
        f"Report {code}{row['report_name']} (country {row['country']}, system {row['system_code']}).{desc} "
        f"Frequency: {row.get('frequency') or 'unspecified'}"
        f"{' (' + row['frequency_other'] + ')' if row.get('frequency_other') else ''}. Layout: {layout}"
        f"{sample_note}"
    )


def _embed(report_id: str) -> None:
    row = db.fetchone("SELECT * FROM kb_reports WHERE report_id = ?", [report_id])
    if row is None:
        return
    vector = llm.embed(text_for(row)[:2000]) if llm.embedding_configured() else None
    db.exec_sql("UPDATE kb_reports SET embedding = ? WHERE report_id = ?", [vector, report_id])


def _payload(row: dict) -> dict:
    return {
        "report_id": row["report_id"],
        "country": row["country"],
        "system_code": row["system_code"],
        "report_code": row.get("report_code") or "",
        "report_name": row["report_name"],
        "frequency": row.get("frequency") or "",
        "frequency_other": row.get("frequency_other") or "",
        "description": row["description"],
        "layout": row["layout"],
        "sample": _sample_public(row.get("sample") or ""),
        "created_at": row["created_at"],
        "last_modified": row["last_modified"],
    }


def list_reports() -> list[dict]:
    rows = db.fetchall(
        """
        SELECT * FROM kb_reports
        ORDER BY country, system_code, report_code, report_name
        """
    )
    return [_payload(row) for row in rows]


def get_report(report_id: str) -> dict:
    row = db.fetchone("SELECT * FROM kb_reports WHERE report_id = ?", [report_id])
    if row is None:
        raise AppError(404, "Report not found")
    return _payload(row)


def create_report(user: dict, body: dict) -> dict:
    _require_editor(user)
    country = _country(str(body.get("country") or ""))
    system_code = _system_code(str(body.get("system_code") or body.get("system") or ""))
    code = _report_code(str(body.get("report_code") or ""))
    name = (body.get("report_name") or "").strip()
    if not name:
        raise AppError(400, "Report name is required.")
    if len(name) > MAX_NAME:
        raise AppError(400, f"Report name must be {MAX_NAME} characters or fewer.")
    description = (body.get("description") or "").strip()
    layout = (body.get("layout") or "").strip()
    if len(description) > MAX_DESCRIPTION:
        raise AppError(400, f"Description must be {MAX_DESCRIPTION} characters or fewer.")
    if len(layout) > MAX_LAYOUT:
        raise AppError(400, f"Layout must be {MAX_LAYOUT} characters or fewer.")
    frequency, frequency_other = _frequency(body)
    _reject_duplicate(country, system_code, code, name)
    report_id = str(uuid.uuid4())
    now = _now()
    db.exec_sql(
        """
        INSERT INTO kb_reports (
            report_id, country, system_code, report_code, report_name, frequency, frequency_other, description, layout, created_at, last_modified
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [report_id, country, system_code, code, name, frequency, frequency_other, description, layout, now, now],
    )
    _embed(report_id)
    gcs.upload_db()
    return get_report(report_id)


def update_report(user: dict, report_id: str, body: dict) -> dict:
    _require_editor(user)
    current = db.fetchone("SELECT * FROM kb_reports WHERE report_id = ?", [report_id])
    if current is None:
        raise AppError(404, "Report not found")
    country = _country(str(body.get("country", current["country"]) or ""))
    system_code = _system_code(
        str(body.get("system_code", body.get("system", current["system_code"])) or ""),
        allow=current["system_code"],
    )
    code = _report_code(str(body.get("report_code", current.get("report_code")) or ""))
    name = (body.get("report_name", current["report_name"]) or "").strip()
    if not name:
        raise AppError(400, "Report name is required.")
    if len(name) > MAX_NAME:
        raise AppError(400, f"Report name must be {MAX_NAME} characters or fewer.")
    description = (body.get("description", current["description"]) or "").strip()
    layout = (body.get("layout", current["layout"]) or "").strip()
    if len(description) > MAX_DESCRIPTION:
        raise AppError(400, f"Description must be {MAX_DESCRIPTION} characters or fewer.")
    if len(layout) > MAX_LAYOUT:
        raise AppError(400, f"Layout must be {MAX_LAYOUT} characters or fewer.")
    frequency, frequency_other = _frequency(body, current)
    _reject_duplicate(country, system_code, code, name, ignore=report_id)
    db.exec_sql(
        """
        UPDATE kb_reports
        SET country = ?, system_code = ?, report_code = ?, report_name = ?, frequency = ?, frequency_other = ?, description = ?, layout = ?, last_modified = ?
        WHERE report_id = ?
        """,
        [country, system_code, code, name, frequency, frequency_other, description, layout, _now(), report_id],
    )
    _embed(report_id)
    gcs.upload_db()
    return get_report(report_id)


def _reject_duplicate(country: str, system_code: str, code: str, name: str, ignore: str | None = None) -> None:
    same_code = db.fetchone(
        """
        SELECT report_id FROM kb_reports
        WHERE country = ? AND system_code = ? AND lower(report_code) = lower(?) AND report_id <> ?
        """,
        [country, system_code, code, ignore or ""],
    )
    if same_code is not None:
        raise AppError(400, f"Report ID {code} already exists for {country} / {system_code}.")
    same_name = db.fetchone(
        """
        SELECT report_id FROM kb_reports
        WHERE country = ? AND system_code = ? AND lower(report_name) = lower(?) AND report_id <> ?
        """,
        [country, system_code, name, ignore or ""],
    )
    if same_name is not None:
        raise AppError(400, f"Report {name} already exists for {country} / {system_code}.")


def delete_report(user: dict, report_id: str) -> None:
    _require_editor(user)
    row = db.fetchone("SELECT sample FROM kb_reports WHERE report_id = ?", [report_id])
    if row is None:
        raise AppError(404, "Report not found")
    _remove_sample_file(report_id, _sample_meta(row.get("sample") or ""))
    db.exec_sql("DELETE FROM kb_reports WHERE report_id = ?", [report_id])
    gcs.upload_db()


def _sample_kind(data: bytes, filename: str) -> str | None:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if data.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    if data.startswith(b"%PDF-"):
        return "pdf"
    if Path(filename or "").suffix.lower() == ".txt" and b"\x00" not in data:
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            return None
        return "txt"
    return None


def _sample_meta(raw: str) -> dict | None:
    try:
        parsed = json.loads(raw or "")
    except json.JSONDecodeError:
        return None
    if not isinstance(parsed, dict):
        return None
    stored = str(parsed.get("stored") or "")
    filename = str(parsed.get("filename") or "")
    media_type = str(parsed.get("media_type") or "")
    if not _STORED_NAME.fullmatch(stored) or not filename or media_type not in SAMPLE_TYPES.values():
        return None
    return {"stored": stored, "filename": filename, "media_type": media_type}


def _sample_public(raw: str) -> dict | None:
    meta = _sample_meta(raw)
    if meta is None:
        return None
    return {"filename": meta["filename"], "media_type": meta["media_type"]}


def _sample_folder(report_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-fA-F-]{8,64}", report_id):
        raise AppError(404, "Report not found")
    return data_dir() / "uploads" / "reports" / report_id


def _remove_sample_file(report_id: str, meta: dict | None) -> None:
    if meta is None:
        return
    path = _sample_folder(report_id) / meta["stored"]
    if path.is_file():
        path.unlink()
    gcs.delete_file(f"uploads/reports/{report_id}/{meta['stored']}")


def save_sample(user: dict, report_id: str, filename: str, data: bytes) -> dict:
    _require_editor(user)
    if db.fetchone("SELECT report_id FROM kb_reports WHERE report_id = ?", [report_id]) is None:
        raise AppError(404, "Report not found")
    if not data:
        raise AppError(400, "The sample is empty.")
    if len(data) > MAX_SAMPLE_BYTES:
        raise AppError(400, "Report samples must be 20 MB or smaller.")
    kind = _sample_kind(data, filename)
    if kind is None:
        raise AppError(400, "Report samples must be a PDF, text, PNG, JPEG, GIF or WEBP file.")
    current = db.fetchone("SELECT sample FROM kb_reports WHERE report_id = ?", [report_id])
    _remove_sample_file(report_id, _sample_meta((current or {}).get("sample") or ""))
    folder = _sample_folder(report_id)
    folder.mkdir(parents=True, exist_ok=True)
    stored = f"{uuid.uuid4().hex}.{kind}"
    path = folder / stored
    path.write_bytes(data)
    display = Path(filename or "sample").name or "sample"
    meta = {"filename": display[:200], "stored": stored, "media_type": SAMPLE_TYPES[kind]}
    db.exec_sql("UPDATE kb_reports SET sample = ?, last_modified = ? WHERE report_id = ?", [json.dumps(meta), _now(), report_id])
    gcs.persist_file(path, f"uploads/reports/{report_id}/{stored}")
    gcs.upload_db()
    return get_report(report_id)


def delete_sample(user: dict, report_id: str) -> dict:
    _require_editor(user)
    row = db.fetchone("SELECT sample FROM kb_reports WHERE report_id = ?", [report_id])
    if row is None:
        raise AppError(404, "Report not found")
    _remove_sample_file(report_id, _sample_meta(row.get("sample") or ""))
    db.exec_sql("UPDATE kb_reports SET sample = '', last_modified = ? WHERE report_id = ?", [_now(), report_id])
    gcs.upload_db()
    return get_report(report_id)


def read_sample(report_id: str) -> tuple[Path, str, str]:
    row = db.fetchone("SELECT sample FROM kb_reports WHERE report_id = ?", [report_id])
    if row is None:
        raise AppError(404, "Report not found")
    meta = _sample_meta(row.get("sample") or "")
    if meta is None:
        raise AppError(404, "Report sample not found")
    path = _sample_folder(report_id) / meta["stored"]
    if not path.is_file():
        gcs.fetch_file(path, f"uploads/reports/{report_id}/{meta['stored']}")
    if not path.is_file():
        raise AppError(404, "Report sample not found")
    return path, meta["media_type"], meta["filename"]


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            if rows and any(h.strip().upper() == "REPORT_ID" for h in rows[0]):
                return rows
        sheet = workbook.worksheets[0]
        return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppError(400, "Could not read the Excel file. Upload an .xlsx file.") from exc


def parse_and_validate(data: bytes) -> list[dict]:
    rows = _read_sheet(data)
    if not rows:
        raise AppError(400, "The Excel file is empty.")
    header = [value.strip().upper() for value in rows[0]]
    missing = [name for name in HEADERS if name not in header]
    if missing:
        raise AppError(400, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(HEADERS)}.")
    index = {name: header.index(name) for name in HEADERS}
    records: list[dict] = []
    errors: list[dict] = []
    seen_codes: dict[str, int] = {}
    seen_names: dict[str, int] = {}

    def error(row_number: int, field: str, message: str) -> None:
        errors.append({"row": row_number, "field": field, "message": message})

    for offset, row in enumerate(rows[1:], start=2):
        values = {name: row[index[name]] if index[name] < len(row) else "" for name in HEADERS}
        if not any(values.values()):
            continue
        code = values["REPORT_ID"].strip()
        name = values["REPORT_NAME"].strip()
        description = values["REPORT_DESCRIPTION"].strip()
        layout = values["REPORT_LAYOUT"].strip()
        frequency = ""
        frequency_other = ""
        try:
            frequency, frequency_other = _frequency_cell(values["FREQUENCY"])
        except AppError as exc:
            error(offset, "FREQUENCY", exc.detail)
        if not code:
            error(offset, "REPORT_ID", "REPORT_ID is required.")
        elif len(code) > MAX_CODE:
            error(offset, "REPORT_ID", f"REPORT_ID must be {MAX_CODE} characters or fewer.")
        code_key = code.upper()
        if code and code_key in seen_codes:
            error(offset, "REPORT_ID", f"Duplicate report ID '{code}' (first seen on row {seen_codes[code_key]}).")
        elif code:
            seen_codes[code_key] = offset
        if not name:
            error(offset, "REPORT_NAME", "REPORT_NAME is required.")
        name_key = name.upper()
        if name and name_key in seen_names:
            error(offset, "REPORT_NAME", f"Duplicate report name '{name}' (first seen on row {seen_names[name_key]}).")
        elif name:
            seen_names[name_key] = offset
        records.append(
            {
                "report_code": code,
                "report_name": name,
                "description": description,
                "frequency": frequency,
                "frequency_other": frequency_other,
                "layout": layout,
            }
        )

    if len(records) > MAX_ROWS:
        raise AppError(400, f"At most {MAX_ROWS} rows can be uploaded at once.")
    if errors:
        raise AppError(400, f"{len(errors)} problem(s) found. Nothing was imported.", errors)
    if not records:
        raise AppError(400, "The Excel file has no data rows.")
    return records


def upload_from_excel(user: dict, country_raw: str, system_raw: str, filename: str, data: bytes) -> dict:
    _require_editor(user)
    country = _country(country_raw)
    system_code = _system_code(system_raw)
    if not (filename or "").lower().endswith(".xlsx"):
        raise AppError(400, "Upload an Excel file (.xlsx).")
    if len(data) > MAX_BYTES:
        raise AppError(400, "The file is larger than 20 MB.")
    records = parse_and_validate(data)
    now = _now()
    for rec in records:
        existing = db.fetchone(
            """
            SELECT report_id, description, layout, frequency, frequency_other FROM kb_reports
            WHERE country = ? AND system_code = ? AND lower(report_code) = lower(?)
            """,
            [country, system_code, rec["report_code"]],
        )
        _reject_duplicate(
            country,
            system_code,
            rec["report_code"],
            rec["report_name"],
            ignore=existing["report_id"] if existing else None,
        )
        if existing is None:
            report_id = str(uuid.uuid4())
            db.exec_sql(
                """
                INSERT INTO kb_reports (
                    report_id, country, system_code, report_code, report_name, frequency, frequency_other, description, layout, created_at, last_modified
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    report_id,
                    country,
                    system_code,
                    rec["report_code"],
                    rec["report_name"],
                    rec["frequency"],
                    rec["frequency_other"],
                    rec["description"],
                    rec["layout"],
                    now,
                    now,
                ],
            )
            _embed(report_id)
        else:
            description = rec["description"] or existing["description"]
            layout = rec["layout"] or existing["layout"]
            frequency = rec["frequency"] or existing["frequency"] or ""
            frequency_other = rec["frequency_other"] if rec["frequency"] else (existing["frequency_other"] or "")
            db.exec_sql(
                """
                UPDATE kb_reports
                SET report_code = ?, report_name = ?, frequency = ?, frequency_other = ?, description = ?, layout = ?, last_modified = ?
                WHERE report_id = ?
                """,
                [
                    rec["report_code"],
                    rec["report_name"],
                    frequency,
                    frequency_other,
                    description,
                    layout,
                    now,
                    existing["report_id"],
                ],
            )
            _embed(existing["report_id"])

    from app.services import reindex

    reindex.embed_missing_in_background()
    gcs.upload_db()
    return {"reports": list_reports(), "uploaded": len(records)}
