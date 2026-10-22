"""Structured knowledge-base capabilities organized by country and system (L1 / L2 / L3)."""

import io
import json
import re
import uuid
from datetime import datetime, timezone

from app import db
from app.config import kb_countries
from app.errors import AppError
from app.permissions import require
from app.services import gcs, llm, systems

HEADERS = ("COUNTRY", "L1_CAPABILITY", "L2_CAPABILITY", "L3_CAPABILITY", "DESCRIPTION")
MAX_ROWS = 5000
MAX_BYTES = 20 * 1024 * 1024
MAX_LEVEL = 200
MAX_DESCRIPTION = 4000
MAX_BEHAVIOR = 20000
MAX_RULES = 20000
# One embedding call can carry this much. Diagrams are placed first so they are not cut off by long rules.
EMBED_CHARS = 8000
MAX_DIAGRAMS = 30
MAX_DIAGRAM_TITLE = 120
MAX_DIAGRAM_TEXT = 20000


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _require_editor(user: dict) -> None:
    require(user, "knowledge", "write")


def _canonical(value: str, allowed: tuple[str, ...]) -> str | None:
    lookup = {item.lower(): item for item in allowed}
    return lookup.get((value or "").strip().lower())


def country_label(value: str) -> str:
    """Empty means the capability applies in every catalog country."""
    codes = [part.strip() for part in (value or "").split(",") if part.strip()]
    return ", ".join(codes) if codes else "All countries"


def _country(value: str) -> str:
    """Blank is all countries. Several codes are stored in catalog order, comma-separated."""
    text = (value or "").strip()
    if not text:
        return ""
    allowed = kb_countries()
    order = {code: index for index, code in enumerate(allowed)}
    chosen: list[str] = []
    seen: set[str] = set()
    for part in re.split(r"[\s,]+", text):
        if not part:
            continue
        canonical = _canonical(part, allowed)
        if canonical is None:
            raise AppError(400, f"Country must be one of: {', '.join(allowed)}.")
        if canonical not in seen:
            seen.add(canonical)
            chosen.append(canonical)
    chosen.sort(key=lambda code: order[code])
    return ",".join(chosen)


def _diagram_list(raw) -> list[dict]:
    if not raw:
        return []
    if isinstance(raw, list):
        data = raw
    else:
        try:
            data = json.loads(raw)
        except (TypeError, json.JSONDecodeError):
            return []
    return [item for item in data if isinstance(item, dict)] if isinstance(data, list) else []


# First word of a Mermaid diagram. Mermaid rejects anything else when the picture is drawn.
MERMAID_TYPES = frozenset(
    {
        "flowchart",
        "graph",
        "sequencediagram",
        "classdiagram",
        "statediagram",
        "statediagram-v2",
        "erdiagram",
        "journey",
        "gantt",
        "pie",
        "quadrantchart",
        "requirementdiagram",
        "gitgraph",
        "mindmap",
        "timeline",
        "zenuml",
        "sankey-beta",
        "xychart-beta",
        "block-beta",
        "packet-beta",
        "kanban",
        "architecture-beta",
        "radar-beta",
        "c4context",
        "c4container",
        "c4component",
        "c4dynamic",
        "c4deployment",
        "treemap",
        "usecase-beta",
        "treeview-beta",
        "swimlane-beta",
    }
)


def sentence_case(value: str) -> str:
    """First letter of the text, and of each following sentence, in capitals. The rest is lower case."""
    text = " ".join((value or "").split())
    if not text:
        return ""
    chars = list(text.lower())
    chars[0] = chars[0].upper()
    index = 0
    while index < len(chars) - 1:
        if chars[index] in ".!?" and chars[index + 1] == " ":
            cursor = index + 1
            while cursor < len(chars) and chars[cursor] == " ":
                cursor += 1
            if cursor < len(chars):
                chars[cursor] = chars[cursor].upper()
            index = cursor
        else:
            index += 1
    return "".join(chars)


def _mermaid_keyword(text: str) -> str | None:
    lines = (text or "").splitlines()
    index = 0
    if lines and lines[0].strip() == "---":
        index = 1
        while index < len(lines) and lines[index].strip() != "---":
            index += 1
        index += 1
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line or line.startswith("%%"):
            continue
        token = re.split(r"[\s;]+", line, maxsplit=1)[0].lower()
        if token in MERMAID_TYPES:
            return token
        return None
    return None


def _diagrams(value) -> str:
    items = value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return "[]"
        try:
            items = json.loads(text)
        except json.JSONDecodeError as exc:
            raise AppError(400, "Diagrams must be a list.") from exc
    if items is None:
        return "[]"
    if not isinstance(items, list):
        raise AppError(400, "Diagrams must be a list.")
    if len(items) > MAX_DIAGRAMS:
        raise AppError(400, f"At most {MAX_DIAGRAMS} diagrams.")
    cleaned = []
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            raise AppError(400, "Each diagram needs a title and Mermaid text.")
        title = str(item.get("title") or "").strip() or f"Diagram {index}"
        mermaid = str(item.get("mermaid") or "").strip()
        if len(title) > MAX_DIAGRAM_TITLE:
            raise AppError(400, f"Diagram titles must be {MAX_DIAGRAM_TITLE} characters or fewer.")
        if not mermaid:
            raise AppError(400, f"Diagram '{title}' is empty.")
        if _mermaid_keyword(mermaid) is None:
            raise AppError(400, f"Diagram '{title}' must be a Mermaid diagram, such as flowchart, sequenceDiagram, swimlane-beta, journey, C4Context, architecture-beta, block-beta, requirementDiagram, usecase-beta, mindmap, or treeView-beta.")
        if len(mermaid) > MAX_DIAGRAM_TEXT:
            raise AppError(400, f"Diagram '{title}' is too long.")
        diagram_id = str(item.get("id") or "").strip() or str(uuid.uuid4())
        cleaned.append({"id": diagram_id, "title": title, "mermaid": mermaid})
    return json.dumps(cleaned)


def _behavior(value: str) -> str:
    text = (value or "").strip()
    if len(text) > MAX_BEHAVIOR:
        raise AppError(400, f"Country specific behavior must be {MAX_BEHAVIOR} characters or fewer.")
    return text


def _system_code(value: str) -> str:
    code = (value or "").strip().upper()
    if not code:
        raise AppError(400, "System is required.")
    if code not in systems.existing_codes():
        raise AppError(400, f"System '{value}' is not a known system code.")
    return code


def _level(value: str, label: str) -> str:
    text = sentence_case(value)
    if not text:
        raise AppError(400, f"{label} is required.")
    if len(text) > MAX_LEVEL:
        raise AppError(400, f"{label} must be {MAX_LEVEL} characters or fewer.")
    return text


def heuristic_description(l1: str, l2: str, l3: str) -> str:
    return f"Provides {l3} as a {l2} feature of {l1}."


def infer_descriptions(levels: list[tuple[str, str, str]]) -> list[str]:
    """One short sentence per L1/L2/L3. Uses the model when it is configured."""
    results = [heuristic_description(l1, l2, l3) for l1, l2, l3 in levels]
    if not levels or not llm.configured():
        return results
    size = 40
    for start in range(0, len(levels), size):
        batch = levels[start : start + size]
        try:
            raw = llm.complete(
                "You write short catalogue descriptions for product capabilities. "
                "Each item has an area (l1), a group (l2) and a feature (l3). "
                "Write one plain sentence that says what the feature lets someone do, using the area and group only as context. "
                "Do not join the three names with '>' or '/'. "
                'Return JSON only: {"descriptions":["..."]} with one string per item, in the same order.',
                json.dumps([{"l1": l1, "l2": l2, "l3": l3} for l1, l2, l3 in batch]),
                0.2,
            )
            start_brace, end_brace = raw.find("{"), raw.rfind("}")
            if start_brace == -1 or end_brace == -1:
                continue
            sentences = json.loads(raw[start_brace : end_brace + 1]).get("descriptions") or []
            for index, sentence in enumerate(sentences):
                text = " ".join(str(sentence).split()).strip().strip('"')
                if text and index < len(batch):
                    results[start + index] = text[:MAX_DESCRIPTION]
        except Exception:
            continue
    return results


def inferred_description(l1: str, l2: str, l3: str) -> str:
    return infer_descriptions([(l1, l2, l3)])[0]


def text_for(row: dict) -> str:
    """Everything the chat should be able to answer from: names, behavior, rules, and the diagram source."""
    parts = [
        f"Capability {row['l1_capability']} / {row['l2_capability']} / {row['l3_capability']} "
        f"(country {country_label(row.get('country') or '')}, system {row['system_code']})."
    ]
    if (row.get("description") or "").strip():
        parts.append(row["description"].strip())
    diagrams = []
    for item in _diagram_list(row.get("sequence_diagrams")):
        title = str(item.get("title") or "Diagram").strip()
        mermaid = str(item.get("mermaid") or "").strip()
        diagrams.append(f"Diagram {title}:\n{mermaid}" if mermaid else f"Diagram {title}")
    if diagrams:
        parts.append("\n".join(diagrams))
    behavior = (row.get("country_behavior") or "").strip()
    if behavior:
        parts.append(f"Country specific behavior: {behavior}")
    rules = (row.get("processing_rules") or "").strip()
    if rules:
        parts.append(f"Processing rules: {rules}")
    return "\n".join(parts)


def _embed(capability_id: str) -> None:
    row = db.fetchone("SELECT * FROM kb_capabilities WHERE capability_id = ?", [capability_id])
    if row is None:
        return
    vector = llm.embed(text_for(row)[:EMBED_CHARS]) if llm.embedding_configured() else None
    db.exec_sql("UPDATE kb_capabilities SET embedding = ? WHERE capability_id = ?", [vector, capability_id])


def _linked(capability_id: str) -> dict:
    interfaces = db.fetchall(
        """
        SELECT i.interface_id, i.channel, i.country, i.interface_name, i.endpoint, i.description, i.protocol, i.source, i.destination
        FROM kb_capability_interfaces link
        JOIN interfaces i ON i.interface_id = link.interface_id
        WHERE link.capability_id = ?
        ORDER BY i.channel, i.interface_name
        """,
        [capability_id],
    )
    tables = db.fetchall(
        """
        SELECT t.table_id, t.country, t.database_name, t.schema_name, t.table_name, t.description
        FROM kb_capability_tables link
        JOIN kb_tables t ON t.table_id = link.table_id
        WHERE link.capability_id = ?
        ORDER BY t.database_name, t.table_name
        """,
        [capability_id],
    )
    batches = db.fetchall(
        """
        SELECT b.batch_id, b.country, b.system_code, b.batch_name, b.batch_type, b.description
        FROM kb_capability_batches link
        JOIN kb_batches b ON b.batch_id = link.batch_id
        WHERE link.capability_id = ?
        ORDER BY b.batch_name
        """,
        [capability_id],
    )
    reports = db.fetchall(
        """
        SELECT r.report_id, r.country, r.system_code, r.report_name, r.description
        FROM kb_capability_reports link
        JOIN kb_reports r ON r.report_id = link.report_id
        WHERE link.capability_id = ?
        ORDER BY r.report_name
        """,
        [capability_id],
    )
    return {
        "interfaces": [
            {
                "interface_id": row["interface_id"],
                "channel": row["channel"],
                "country": row["country"],
                "interface_name": row["interface_name"],
                "endpoint": row.get("endpoint") or "",
                "description": row["description"],
                "protocol": row["protocol"],
                "source": row["source"],
                "destination": row["destination"],
            }
            for row in interfaces
        ],
        "tables": [
            {
                "table_id": row["table_id"],
                "country": row["country"],
                "database_name": row["database_name"],
                "schema_name": row["schema_name"],
                "table_name": row["table_name"],
                "description": row.get("description") or "",
            }
            for row in tables
        ],
        "batches": [
            {
                "batch_id": row["batch_id"],
                "country": row["country"],
                "system_code": row["system_code"],
                "batch_name": row["batch_name"],
                "batch_type": row["batch_type"],
                "description": row["description"],
            }
            for row in batches
        ],
        "reports": [
            {
                "report_id": row["report_id"],
                "country": row["country"],
                "system_code": row["system_code"],
                "report_name": row["report_name"],
                "description": row["description"],
            }
            for row in reports
        ],
    }


def _payload(row: dict, *, with_links: bool = False) -> dict:
    body = {
        "capability_id": row["capability_id"],
        "country": row["country"],
        "system_code": row["system_code"],
        "l1_capability": sentence_case(row["l1_capability"]),
        "l2_capability": sentence_case(row["l2_capability"]),
        "l3_capability": sentence_case(row["l3_capability"]),
        "description": sentence_case(row["description"] or ""),
        "sequence_diagrams": _diagram_list(row.get("sequence_diagrams")),
        "processing_rules": row.get("processing_rules") or "",
        "country_behavior": row.get("country_behavior") or "",
        "created_at": row["created_at"],
        "last_modified": row["last_modified"],
    }
    if with_links:
        body.update(_linked(row["capability_id"]))
    return body


def list_capabilities() -> list[dict]:
    rows = db.fetchall(
        """
        SELECT * FROM kb_capabilities
        ORDER BY country, system_code, l1_capability, l2_capability, l3_capability
        """
    )
    return [_payload(row) for row in rows]


def markdown_body(record: dict) -> tuple[str, str]:
    """Title and Markdown for the capability download. Diagrams stay as Mermaid fences."""
    title = " / ".join(part for part in (record.get("l1_capability"), record.get("l2_capability"), record.get("l3_capability")) if part) or "Capability"
    country = record.get("country") or "All countries"
    lines = [
        f"**System:** {record.get('system_code') or ''}",
        f"**Applicable for:** {country}",
        "",
        "## Description",
        "",
        record.get("description") or "",
        "",
        "## Country specific behavior",
        "",
        record.get("country_behavior") or "",
        "",
        "## Diagrams",
        "",
    ]
    diagrams = record.get("sequence_diagrams") or []
    if not diagrams:
        lines.extend(["No diagrams.", ""])
    for diagram in diagrams:
        lines.extend([f"### {diagram.get('title') or 'Diagram'}", "", "```mermaid", diagram.get("mermaid") or "", "```", ""])
    lines.extend(["## Processing rules", "", record.get("processing_rules") or "", ""])
    for heading, key, label in (
        ("Interfaces", "interfaces", "interface_name"),
        ("Tables", "tables", "table_name"),
        ("Batches", "batches", "batch_name"),
        ("Reports", "reports", "report_name"),
    ):
        lines.extend([f"## {heading}", ""])
        rows = record.get(key) or []
        if not rows:
            lines.append("None.")
        else:
            lines.extend(f"- {row.get(label) or ''}" for row in rows)
        lines.append("")
    return title, "\n".join(lines).rstrip() + "\n"


def get_capability(capability_id: str) -> dict:
    row = db.fetchone("SELECT * FROM kb_capabilities WHERE capability_id = ?", [capability_id])
    if row is None:
        raise AppError(404, "Capability not found")
    return _payload(row, with_links=True)


def _validate_fields(l1: str, l2: str, l3: str, description: str, rules: str) -> tuple[str, str, str, str, str]:
    l1 = _level(l1, "L1 capability")
    l2 = _level(l2, "L2 capability")
    l3 = _level(l3, "L3 capability")
    description = sentence_case(description) or inferred_description(l1, l2, l3)
    description = sentence_case(description)
    rules = (rules or "").strip()
    if len(description) > MAX_DESCRIPTION:
        raise AppError(400, f"Description must be {MAX_DESCRIPTION} characters or fewer.")
    if len(rules) > MAX_RULES:
        raise AppError(400, f"Processing rules must be {MAX_RULES} characters or fewer.")
    return l1, l2, l3, description, rules


def create_capability(user: dict, body: dict) -> dict:
    _require_editor(user)
    country = _country(str(body.get("country") or ""))
    behavior = _behavior(str(body.get("country_behavior") or ""))
    system_code = _system_code(str(body.get("system_code") or body.get("system") or ""))
    l1, l2, l3, description, rules = _validate_fields(
        str(body.get("l1_capability") or ""),
        str(body.get("l2_capability") or ""),
        str(body.get("l3_capability") or ""),
        str(body.get("description") or ""),
        str(body.get("processing_rules") or ""),
    )
    diagrams = _diagrams(body.get("sequence_diagrams") or [])
    existing = db.fetchone(
        """
        SELECT capability_id FROM kb_capabilities
        WHERE country = ? AND system_code = ?
          AND lower(l1_capability) = lower(?) AND lower(l2_capability) = lower(?) AND lower(l3_capability) = lower(?)
        """,
        [country, system_code, l1, l2, l3],
    )
    if existing is not None:
        raise AppError(400, f"Capability {l1} / {l2} / {l3} already exists for {country_label(country)} / {system_code}.")
    capability_id = str(uuid.uuid4())
    now = _now()
    db.exec_sql(
        """
        INSERT INTO kb_capabilities (
            capability_id, country, system_code, l1_capability, l2_capability, l3_capability,
            description, sequence_diagrams, processing_rules, country_behavior, created_at, last_modified
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [capability_id, country, system_code, l1, l2, l3, description, diagrams, rules, behavior, now, now],
    )
    _embed(capability_id)
    gcs.upload_db()
    return get_capability(capability_id)


def update_capability(user: dict, capability_id: str, body: dict) -> dict:
    _require_editor(user)
    current = db.fetchone("SELECT * FROM kb_capabilities WHERE capability_id = ?", [capability_id])
    if current is None:
        raise AppError(404, "Capability not found")
    country = _country(str(body.get("country", current["country"]) or ""))
    behavior = _behavior(str(body.get("country_behavior", current.get("country_behavior") or "") or ""))
    system_code = _system_code(str(body.get("system_code", body.get("system", current["system_code"])) or ""))
    l1, l2, l3, description, rules = _validate_fields(
        str(body.get("l1_capability", current["l1_capability"]) or ""),
        str(body.get("l2_capability", current["l2_capability"]) or ""),
        str(body.get("l3_capability", current["l3_capability"]) or ""),
        str(body.get("description", current["description"]) or ""),
        str(body.get("processing_rules", current.get("processing_rules") or "") or ""),
    )
    diagrams = _diagrams(body["sequence_diagrams"]) if "sequence_diagrams" in body else (current.get("sequence_diagrams") or "[]")
    if not isinstance(diagrams, str):
        diagrams = _diagrams(diagrams)
    conflict = db.fetchone(
        """
        SELECT capability_id FROM kb_capabilities
        WHERE country = ? AND system_code = ?
          AND lower(l1_capability) = lower(?) AND lower(l2_capability) = lower(?) AND lower(l3_capability) = lower(?)
          AND capability_id <> ?
        """,
        [country, system_code, l1, l2, l3, capability_id],
    )
    if conflict is not None:
        raise AppError(400, f"Capability {l1} / {l2} / {l3} already exists for {country_label(country)} / {system_code}.")
    db.exec_sql(
        """
        UPDATE kb_capabilities
        SET country = ?, system_code = ?, l1_capability = ?, l2_capability = ?, l3_capability = ?,
            description = ?, sequence_diagrams = ?, processing_rules = ?, country_behavior = ?, last_modified = ?
        WHERE capability_id = ?
        """,
        [country, system_code, l1, l2, l3, description, diagrams, rules, behavior, _now(), capability_id],
    )
    _embed(capability_id)
    gcs.upload_db()
    return get_capability(capability_id)


def _clear_links(capability_id: str) -> None:
    for table in (
        "kb_capability_interfaces",
        "kb_capability_tables",
        "kb_capability_batches",
        "kb_capability_reports",
    ):
        db.exec_sql(f"DELETE FROM {table} WHERE capability_id = ?", [capability_id])


def delete_capability(user: dict, capability_id: str) -> None:
    _require_editor(user)
    get_capability(capability_id)
    _clear_links(capability_id)
    db.exec_sql("DELETE FROM kb_capabilities WHERE capability_id = ?", [capability_id])
    gcs.upload_db()


def _require_capability(capability_id: str) -> dict:
    row = db.fetchone("SELECT * FROM kb_capabilities WHERE capability_id = ?", [capability_id])
    if row is None:
        raise AppError(404, "Capability not found")
    return row


def _country_codes(value: str) -> set[str]:
    return {
        part.strip().lower()
        for part in (value or "").split(",")
        if part.strip() and part.strip().lower() not in {"all", "all countries"}
    }


def _require_same_country(capability: dict, row: dict | None, label: str) -> None:
    if row is None:
        raise AppError(404, f"{label} not found")
    allowed = _country_codes(capability.get("country") or "")
    if not allowed:
        return
    item = _country_codes(row.get("country") or "")
    # An item for every country overlaps any capability. Otherwise one shared code is enough.
    if not item or allowed & item:
        return
    raise AppError(
        400,
        f"{label} must be in the same country as the capability ({country_label(capability.get('country') or '')}).",
    )


def link_interface(user: dict, capability_id: str, interface_id: str) -> dict:
    _require_editor(user)
    capability = _require_capability(capability_id)
    _require_same_country(
        capability,
        db.fetchone("SELECT interface_id, country FROM interfaces WHERE interface_id = ?", [interface_id]),
        "Interface",
    )
    if db.fetchone(
        "SELECT 1 FROM kb_capability_interfaces WHERE capability_id = ? AND interface_id = ?",
        [capability_id, interface_id],
    ):
        raise AppError(400, "Interface is already linked to this capability.")
    db.exec_sql(
        "INSERT INTO kb_capability_interfaces (capability_id, interface_id) VALUES (?, ?)",
        [capability_id, interface_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def unlink_interface(user: dict, capability_id: str, interface_id: str) -> dict:
    _require_editor(user)
    _require_capability(capability_id)
    db.exec_sql(
        "DELETE FROM kb_capability_interfaces WHERE capability_id = ? AND interface_id = ?",
        [capability_id, interface_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def link_table(user: dict, capability_id: str, table_id: str) -> dict:
    _require_editor(user)
    capability = _require_capability(capability_id)
    _require_same_country(
        capability,
        db.fetchone("SELECT table_id, country FROM kb_tables WHERE table_id = ?", [table_id]),
        "Table",
    )
    if db.fetchone(
        "SELECT 1 FROM kb_capability_tables WHERE capability_id = ? AND table_id = ?",
        [capability_id, table_id],
    ):
        raise AppError(400, "Table is already linked to this capability.")
    db.exec_sql(
        "INSERT INTO kb_capability_tables (capability_id, table_id) VALUES (?, ?)",
        [capability_id, table_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def unlink_table(user: dict, capability_id: str, table_id: str) -> dict:
    _require_editor(user)
    _require_capability(capability_id)
    db.exec_sql(
        "DELETE FROM kb_capability_tables WHERE capability_id = ? AND table_id = ?",
        [capability_id, table_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def link_batch(user: dict, capability_id: str, batch_id: str) -> dict:
    _require_editor(user)
    capability = _require_capability(capability_id)
    _require_same_country(
        capability,
        db.fetchone("SELECT batch_id, country FROM kb_batches WHERE batch_id = ?", [batch_id]),
        "Batch",
    )
    if db.fetchone(
        "SELECT 1 FROM kb_capability_batches WHERE capability_id = ? AND batch_id = ?",
        [capability_id, batch_id],
    ):
        raise AppError(400, "Batch is already linked to this capability.")
    db.exec_sql(
        "INSERT INTO kb_capability_batches (capability_id, batch_id) VALUES (?, ?)",
        [capability_id, batch_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def unlink_batch(user: dict, capability_id: str, batch_id: str) -> dict:
    _require_editor(user)
    _require_capability(capability_id)
    db.exec_sql(
        "DELETE FROM kb_capability_batches WHERE capability_id = ? AND batch_id = ?",
        [capability_id, batch_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def link_report(user: dict, capability_id: str, report_id: str) -> dict:
    _require_editor(user)
    capability = _require_capability(capability_id)
    _require_same_country(
        capability,
        db.fetchone("SELECT report_id, country FROM kb_reports WHERE report_id = ?", [report_id]),
        "Report",
    )
    if db.fetchone(
        "SELECT 1 FROM kb_capability_reports WHERE capability_id = ? AND report_id = ?",
        [capability_id, report_id],
    ):
        raise AppError(400, "Report is already linked to this capability.")
    db.exec_sql(
        "INSERT INTO kb_capability_reports (capability_id, report_id) VALUES (?, ?)",
        [capability_id, report_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def unlink_report(user: dict, capability_id: str, report_id: str) -> dict:
    _require_editor(user)
    _require_capability(capability_id)
    db.exec_sql(
        "DELETE FROM kb_capability_reports WHERE capability_id = ? AND report_id = ?",
        [capability_id, report_id],
    )
    gcs.upload_db()
    return get_capability(capability_id)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_sheet(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            if rows and any(h.strip().upper() == "L1_CAPABILITY" for h in rows[0]):
                return rows
        sheet = workbook.worksheets[0]
        return [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise AppError(400, "Could not read the Excel file. Upload an .xlsx file.") from exc


def parse_and_validate(data: bytes) -> list[dict]:
    rows = _read_sheet(data)
    if not rows:
        raise AppError(400, "The Excel file is empty.")
    header = [value.strip().upper() for value in rows[0]]
    # Accept friendly aliases for the Excel headers.
    aliases = {
        "L1 CAPABILITY": "L1_CAPABILITY",
        "L2 CAPABILITY": "L2_CAPABILITY",
        "L3 CAPABILITY": "L3_CAPABILITY",
    }
    header = [aliases.get(name, name.replace(" ", "_")) for name in header]
    missing = [name for name in HEADERS if name not in header]
    if missing:
        raise AppError(400, f"Missing column(s): {', '.join(missing)}. Expected: {', '.join(HEADERS)}.")
    index = {name: header.index(name) for name in HEADERS}
    records: list[dict] = []
    errors: list[dict] = []
    seen: dict[str, int] = {}

    def error(row_number: int, field: str, message: str) -> None:
        errors.append({"row": row_number, "field": field, "message": message})

    for offset, row in enumerate(rows[1:], start=2):
        values = {name: row[index[name]] if index[name] < len(row) else "" for name in HEADERS}
        if not any(values.values()):
            continue
        l1 = sentence_case(values["L1_CAPABILITY"])
        l2 = sentence_case(values["L2_CAPABILITY"])
        l3 = sentence_case(values["L3_CAPABILITY"])
        description = sentence_case(values["DESCRIPTION"])
        country = ""
        country_ok = True
        try:
            country = _country(values["COUNTRY"])
        except AppError as exc:
            country_ok = False
            error(offset, "COUNTRY", exc.detail)
        if not l1:
            error(offset, "L1_CAPABILITY", "L1_CAPABILITY is required.")
        if not l2:
            error(offset, "L2_CAPABILITY", "L2_CAPABILITY is required.")
        if not l3:
            error(offset, "L3_CAPABILITY", "L3_CAPABILITY is required.")
        key = f"{country}\0{l1}\0{l2}\0{l3}".upper()
        if country_ok and l1 and l2 and l3 and key in seen:
            error(offset, "L1_CAPABILITY", f"Duplicate COUNTRY/L1/L2/L3 (first seen on row {seen[key]}).")
        elif country_ok and l1 and l2 and l3:
            seen[key] = offset
        if len(l1) > MAX_LEVEL:
            error(offset, "L1_CAPABILITY", f"L1_CAPABILITY must be {MAX_LEVEL} characters or fewer.")
        if len(l2) > MAX_LEVEL:
            error(offset, "L2_CAPABILITY", f"L2_CAPABILITY must be {MAX_LEVEL} characters or fewer.")
        if len(l3) > MAX_LEVEL:
            error(offset, "L3_CAPABILITY", f"L3_CAPABILITY must be {MAX_LEVEL} characters or fewer.")
        if len(description) > MAX_DESCRIPTION:
            error(offset, "DESCRIPTION", f"DESCRIPTION must be {MAX_DESCRIPTION} characters or fewer.")
        records.append(
            {
                "country": country,
                "l1_capability": l1,
                "l2_capability": l2,
                "l3_capability": l3,
                "description": description,
            }
        )

    if len(records) > MAX_ROWS:
        raise AppError(400, f"At most {MAX_ROWS} rows can be uploaded at once.")
    if errors:
        raise AppError(400, f"{len(errors)} problem(s) found. Nothing was imported.", errors)
    if not records:
        raise AppError(400, "The Excel file has no data rows.")
    missing = [index for index, rec in enumerate(records) if not rec["description"] and rec["l1_capability"] and rec["l2_capability"] and rec["l3_capability"]]
    if missing:
        generated = infer_descriptions(
            [(records[index]["l1_capability"], records[index]["l2_capability"], records[index]["l3_capability"]) for index in missing]
        )
        for index, text in zip(missing, generated):
            records[index]["description"] = sentence_case(text)
    return records


def upload_from_excel(user: dict, system_raw: str, filename: str, data: bytes) -> dict:
    _require_editor(user)
    system_code = _system_code(system_raw)
    if not (filename or "").lower().endswith(".xlsx"):
        raise AppError(400, "Upload an Excel file (.xlsx).")
    if len(data) > MAX_BYTES:
        raise AppError(400, "The file is larger than 20 MB.")
    records = parse_and_validate(data)
    now = _now()
    for rec in records:
        existing = db.fetchone(
            """
            SELECT capability_id, description, processing_rules FROM kb_capabilities
            WHERE country = ? AND system_code = ?
              AND lower(l1_capability) = lower(?) AND lower(l2_capability) = lower(?) AND lower(l3_capability) = lower(?)
            """,
            [rec["country"], system_code, rec["l1_capability"], rec["l2_capability"], rec["l3_capability"]],
        )
        if existing is None:
            capability_id = str(uuid.uuid4())
            db.exec_sql(
                """
                INSERT INTO kb_capabilities (
                    capability_id, country, system_code, l1_capability, l2_capability, l3_capability,
                    description, processing_rules, created_at, last_modified
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                [
                    capability_id,
                    rec["country"],
                    system_code,
                    rec["l1_capability"],
                    rec["l2_capability"],
                    rec["l3_capability"],
                    rec["description"],
                    "",
                    now,
                    now,
                ],
            )
            _embed(capability_id)
        else:
            description = rec["description"] or existing["description"]
            db.exec_sql(
                """
                UPDATE kb_capabilities
                SET l1_capability = ?, l2_capability = ?, l3_capability = ?, description = ?, last_modified = ?
                WHERE capability_id = ?
                """,
                [
                    rec["l1_capability"],
                    rec["l2_capability"],
                    rec["l3_capability"],
                    description,
                    now,
                    existing["capability_id"],
                ],
            )
            _embed(existing["capability_id"])

    from app.services import reindex

    reindex.embed_missing_in_background()
    gcs.upload_db()
    return {"capabilities": list_capabilities(), "uploaded": len(records)}
